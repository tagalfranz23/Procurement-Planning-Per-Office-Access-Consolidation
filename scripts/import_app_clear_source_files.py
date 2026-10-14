"""Import populated APP source data and audit the blank APP-CSE template.

The all-funds workbook is stored as a separate read-only APP_NON_CSE record.
Missing 12-column APP fields remain NULL. The zero-quantity APP-CSE catalogue is
validated but is never inserted as procurement requirements.
"""

from argparse import ArgumentParser
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from hashlib import sha256
from pathlib import Path
import re
import sys

from openpyxl import load_workbook


PPMP_ROOT = Path(__file__).resolve().parents[1]
if str(PPMP_ROOT) not in sys.path:
    sys.path.insert(0, str(PPMP_ROOT))

from app import app  # noqa: E402
from models import (  # noqa: E402
    APPItem, APPSection, AnnualProcurementPlan, WorkflowAuditLog, db,
)


REFERENCE_DIR = PPMP_ROOT / 'reference_data' / 'app_2026'
ALL_FUNDS_FILE = 'APP-NON-CSE-FOR-MPHILGEPS-ALL-FUNDS-2026.xlsx'
CSE_FILE = 'APP-CSE 2026 Form.xlsx'
OFFICE = 'ISU Cauayan Campus'
FISCAL_YEAR = 2026
START = date(2026, 1, 1)
END = date(2026, 12, 1)
GROUPS = (
    ('GAA FUND', 6, 46, Decimal('93409200.00')),
    ('FUND 164', 48, 90, Decimal('84147871.30')),
    ('FUND 161', 92, 112, Decimal('8462750.00')),
)
EXPECTED_TOTAL = Decimal('186019821.30')


def money(value):
    return Decimal(str(value or 0)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def normalized_text(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def file_hash(path):
    digest = sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def read_all_funds():
    path = REFERENCE_DIR / ALL_FUNDS_FILE
    if not path.is_file():
        raise RuntimeError(f'Missing source workbook: {ALL_FUNDS_FILE}')
    sheet = load_workbook(path, data_only=False, read_only=False)['APP_101']
    sections = []
    for section_name, first_row, last_row, expected_total in GROUPS:
        rows = []
        for row_number in range(first_row, last_row + 1):
            title = normalized_text(sheet.cell(row_number, 2).value)
            if not title:
                continue
            schedule = {
                normalized_text(sheet.cell(row_number, column).value)
                for column in (6, 7, 8, 9)
            }
            if schedule != {'January-December 2026'}:
                raise RuntimeError(f'Unexpected schedule in {ALL_FUNDS_FILE} row {row_number}.')
            budget = money(sheet.cell(row_number, 12).value) + money(sheet.cell(row_number, 13).value)
            if budget < 0:
                raise RuntimeError(f'Negative budget in {ALL_FUNDS_FILE} row {row_number}.')
            rows.append({
                'source_row': row_number,
                'project_title': title,
                'implementing_unit': normalized_text(sheet.cell(row_number, 4).value),
                # These columns do not exist in this mPhilGEPS workbook.
                'general_description': None,
                'early_procurement_activity': None,
                'bid_evaluation_criteria': None,
                'mode_of_procurement': normalized_text(sheet.cell(row_number, 5).value),
                'procurement_start': START,
                'procurement_end': END,
                'source_of_fund': normalized_text(sheet.cell(row_number, 10).value),
                'estimated_budget': budget,
                'procurement_strategy_tools': None,
                'remarks': normalized_text(sheet.cell(row_number, 14).value) or None,
            })
        actual_total = sum((row['estimated_budget'] for row in rows), Decimal('0.00'))
        if actual_total != expected_total:
            raise RuntimeError(
                f'{section_name} total is {actual_total}, expected {expected_total}; import stopped.'
            )
        sections.append((section_name, rows))

    row_count = sum(len(rows) for _, rows in sections)
    overall = sum(
        (row['estimated_budget'] for _, rows in sections for row in rows), Decimal('0.00')
    )
    if row_count != 105 or overall != EXPECTED_TOTAL:
        raise RuntimeError('All-funds workbook row count or total does not match the verified source.')
    return path, sections


def audit_cse_template():
    path = REFERENCE_DIR / CSE_FILE
    if not path.is_file():
        raise RuntimeError(f'Missing source workbook: {CSE_FILE}')
    formula_sheet = load_workbook(path, data_only=False, read_only=False)['APP-CSE 2026 FORM']
    value_sheet = load_workbook(path, data_only=True, read_only=False)['APP-CSE 2026 FORM']
    item_rows = [
        row for row in range(1, formula_sheet.max_row + 1)
        if formula_sheet.cell(row, 2).value
        and formula_sheet.cell(row, 3).value
        and formula_sheet.cell(row, 4).value
    ]
    month_columns = (5, 6, 7, 10, 11, 12, 15, 16, 17, 20, 21, 22)
    active_rows = [
        row for row in item_rows
        if any(money(formula_sheet.cell(row, column).value) != 0 for column in month_columns)
    ]
    cached_totals = {
        'part_i': money(value_sheet['Y383'].value),
        'part_ii': money(value_sheet['Y391'].value),
    }
    if len(item_rows) != 318 or active_rows or any(cached_totals.values()):
        raise RuntimeError('APP-CSE workbook is not the verified blank zero-quantity template.')
    return {
        'file': CSE_FILE,
        'sha256': file_hash(path),
        'catalog_rows': len(item_rows),
        'active_requirement_rows': 0,
        'cached_total': 0,
        'imported_as_app_cse': False,
        'reason': 'No agency quantities or procurement amounts are entered.',
    }


def source_details(path, sections, cse_audit):
    group_totals = {
        name: float(sum((row['estimated_budget'] for row in rows), Decimal('0.00')))
        for name, rows in sections
    }
    return {
        'description': 'FY 2026 mPhilGEPS APP Non-CSE all-funds workbook',
        'source_label': 'APP Non-CSE all-funds Excel workbook',
        'source_file': path.name,
        'source_sha256': file_hash(path),
        'source_sheet': 'APP_101',
        'source_row_groups': {
            'GAA FUND': '6-46', 'FUND 164': '48-90', 'FUND 161': '92-112',
        },
        'reported_totals': {'overall': float(EXPECTED_TOTAL), **group_totals},
        'reported_variance': 0,
        'unavailable_app_columns': [
            'general_description', 'early_procurement_activity',
            'bid_evaluation_criteria', 'procurement_strategy_tools',
        ],
        'cse_template_audit': cse_audit,
        'workflow_note': 'Source status and approval signatures are absent; record remains Imported.',
    }


def existing_is_identical(plan, sections):
    expected = [row for _, rows in sections for row in rows]
    return (
        plan.source_type == 'REFERENCE_DOCUMENTS'
        and plan.status == 'Imported'
        and len(plan.sections) == len(sections)
        and len(plan.items) == len(expected)
        and [row.project_title for row in plan.items]
        == [row['project_title'] for row in expected]
        and [row.estimated_budget for row in plan.items]
        == [row['estimated_budget'] for row in expected]
    )


def import_all_funds():
    path, sections = read_all_funds()
    cse_audit = audit_cse_template()
    existing = AnnualProcurementPlan.query.filter_by(
        office=OFFICE, fiscal_year=FISCAL_YEAR, plan_type='APP_NON_CSE',
    ).one_or_none()
    if existing is not None:
        if existing_is_identical(existing, sections):
            return existing, False, cse_audit
        raise RuntimeError('A conflicting APP Non-CSE record exists; no data was changed.')

    plan = AnnualProcurementPlan(
        office=OFFICE,
        fiscal_year=FISCAL_YEAR,
        plan_type='APP_NON_CSE',
        consolidation_batch_id=None,
        source_type='REFERENCE_DOCUMENTS',
        source_details=source_details(path, sections, cse_audit),
        status='Imported',
        created_by=None,
        remarks=(
            'Imported directly from the FY 2026 mPhilGEPS all-funds workbook. '
            'Fields absent from that format remain blank. The blank APP-CSE template '
            'was audited but its zero-quantity catalogue was not imported as requirements.'
        ),
    )
    db.session.add(plan)
    db.session.flush()
    for order, (section_name, rows) in enumerate(sections, 1):
        section = APPSection(
            app_id=plan.id, name=section_name, sort_order=order * 10, created_by=None,
        )
        db.session.add(section)
        db.session.flush()
        for row in rows:
            values = {key: value for key, value in row.items() if key != 'source_row'}
            db.session.add(APPItem(
                app_id=plan.id, section_id=section.id,
                source_ppmp_item_id=None, created_by=None, **values,
            ))
    db.session.add(WorkflowAuditLog(
        entity_type='app', entity_id=plan.id,
        action='CLEAR_SOURCE_IMPORTED', performed_by=None,
        details={
            'source_file': path.name,
            'row_count': 105,
            'total': float(EXPECTED_TOTAL),
            'sections': [name for name, _ in sections],
            'cse_active_rows': 0,
        },
    ))
    db.session.commit()
    return plan, True, cse_audit


def main():
    parser = ArgumentParser(description=__doc__)
    action = parser.add_mutually_exclusive_group(required=True)
    action.add_argument('--check-only', action='store_true')
    action.add_argument('--apply', action='store_true')
    args = parser.parse_args()

    path, sections = read_all_funds()
    cse_audit = audit_cse_template()
    summary = {
        'source': path.name,
        'sections': {name: len(rows) for name, rows in sections},
        'row_count': sum(len(rows) for _, rows in sections),
        'total': str(EXPECTED_TOTAL),
        'cse': cse_audit,
    }
    if args.check_only:
        print(summary)
        return

    with app.app_context():
        try:
            plan, created, cse_audit = import_all_funds()
            print({
                'app_id': plan.id, 'created': created, 'status': plan.status,
                'rows': len(plan.items), 'total': str(plan.totals()['overall']),
                'cse_active_rows': cse_audit['active_requirement_rows'],
            })
        except Exception:
            db.session.rollback()
            raise


if __name__ == '__main__':
    main()
