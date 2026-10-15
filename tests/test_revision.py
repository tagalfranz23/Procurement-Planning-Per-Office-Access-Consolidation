import io
import os
import re
import unittest
from datetime import datetime
from unittest.mock import MagicMock, patch

os.environ['DATABASE_URL'] = 'sqlite:///:memory:'
os.environ['FLASK_SECRET_KEY'] = 'test-only-secret'

from sqlalchemy.exc import IntegrityError
from werkzeug.security import generate_password_hash
from openpyxl import load_workbook
from PyPDF2 import PdfReader

import app as application
from models import (
    APPComment, APPItem, APPSection, AnnualProcurementPlan, College,
    ConsolidationBatch, ConsolidationSource, PasswordResetRequest, PPMPItem,
    PPMPNotification, PPMPOrder, PPMPOrderItem, PPMPRecord, PPMPRecordItem, PPMPReport,
    User, UserActivityLog, WishlistItem, WishlistRequest, db,
)


class RevisionRulesTest(unittest.TestCase):
    def setUp(self):
        application.app.config.update(TESTING=True)
        self.context = application.app.app_context()
        self.context.push()
        db.drop_all()
        db.create_all()
        application._RESET_RATE_BUCKETS.clear()
        db.session.add(User(username='admin', full_name='Admin', email='admin@test.local',
                            password_hash=generate_password_hash('password123'), role='ADMIN',
                            office='ADMINISTRATION', account_status='ACTIVE', is_active=True))
        db.session.commit()
        self.client = application.app.test_client()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def registration(self, **changes):
        payload = {'first_name': 'Procurement', 'last_name': 'Officer', 'email': 'procurement@isu.edu.ph',
                   'username': 'procurement', 'password': 'password123', 'role': 'PROCUREMENT',
                   'password_confirmation': 'password123', 'office': 'PROCUREMENT'}
        payload.update(changes)
        return self.client.post('/api/auth/register', json=payload)

    def test_reference_image_app_transcription_and_provenance(self):
        from scripts.import_app_reference_images import (
            REPORTED_TOTALS, validate_transcription,
        )

        summary = validate_transcription()
        self.assertEqual(summary['row_count'], 40)
        self.assertEqual(summary['section_count'], 4)
        self.assertEqual(summary['row_total'], 93_459_200)
        self.assertEqual(summary['reported_variance'], 0)
        self.assertEqual(REPORTED_TOTALS['overall'], 93_459_200)

        plan = AnnualProcurementPlan(
            office='ISU Cauayan Campus', fiscal_year=2026, plan_type='APP',
            consolidation_batch_id=None, source_type='REFERENCE_DOCUMENTS',
            source_details={'image_count': 4}, status='Finalized',
        )
        db.session.add(plan)
        db.session.commit()
        detail = application.app_detail_dict(plan)
        self.assertEqual(detail['source']['source_type'], 'REFERENCE_DOCUMENTS')
        self.assertEqual(detail['source']['source_details']['image_count'], 4)
        self.assertEqual(detail['source']['consolidated_item_count'], 0)
        self.assertIsNone(detail['expected_reviewer'])

    def test_ppmp_image_wishlist_import_is_validated_and_idempotent(self):
        from scripts.import_ppmp_wishlist_images import import_rows, validate_transcription

        summary = validate_transcription()
        self.assertEqual(summary['importable_rows'], 29)
        self.assertEqual(summary['quantified_row_total'], '225778.00')
        self.assertEqual(summary['printed_vs_transcribed_variance'], '1100.00')
        rows, changed, _ = import_rows()
        self.assertTrue(changed)
        self.assertEqual(len(rows), 29)
        self.assertEqual(WishlistItem.query.filter_by(status='Available').count(), 29)
        rows, changed, _ = import_rows()
        self.assertFalse(changed)
        self.assertEqual(len(rows), 29)
        self.assertEqual(PPMPNotification.query.count(), 0)

    def test_clear_all_funds_and_blank_cse_sources_are_validated(self):
        from scripts.import_app_clear_source_files import (
            EXPECTED_TOTAL, audit_cse_template, read_all_funds,
        )

        _, sections = read_all_funds()
        self.assertEqual([name for name, _ in sections], ['GAA FUND', 'FUND 164', 'FUND 161'])
        self.assertEqual(sum(len(rows) for _, rows in sections), 105)
        self.assertEqual(
            sum((row['estimated_budget'] for _, rows in sections for row in rows), 0),
            EXPECTED_TOTAL,
        )
        cse = audit_cse_template()
        self.assertEqual(cse['catalog_rows'], 318)
        self.assertEqual(cse['active_requirement_rows'], 0)
        self.assertFalse(cse['imported_as_app_cse'])

    def test_pending_registration_reserves_role(self):
        self.assertEqual(self.registration().status_code, 201)
        response = self.registration(email='second@isu.edu.ph', username='second')
        self.assertEqual(response.status_code, 409)
        self.assertIn('role', response.get_json()['error'].lower())

    def test_duplicate_email_is_blocked(self):
        self.assertEqual(self.registration().status_code, 201)
        response = self.registration(role='EO', office='EO', username='eo')
        self.assertEqual(response.status_code, 409)
        self.assertIn('email', response.get_json()['error'].lower())

    def test_forgot_password_does_not_create_an_admin_recovery_queue(self):
        user = User(username='recover_me', full_name='Recover Me', email='recover@isu.edu.ph',
                    password_hash=generate_password_hash('old-password'), role='CBAO', office='CBAO',
                    account_status='ACTIVE', is_active=True)
        db.session.add(user)
        db.session.commit()
        recovery_client = application.app.test_client()
        login_html = recovery_client.get('/').get_data(as_text=True)
        self.assertIn('Forgot Password?', login_html)
        self.assertIn('id="forgot-view"', login_html)
        self.assertIn('id="reset-view"', login_html)
        self.assertIn('Email delivery is still being configured', login_html)

        known = recovery_client.post('/api/auth/forgot-password', json={'email': user.email})
        unknown = recovery_client.post('/api/auth/forgot-password', json={'email': 'unknown@isu.edu.ph'})
        self.assertEqual(known.status_code, 200)
        self.assertEqual(known.get_json()['message'], unknown.get_json()['message'])
        self.assertEqual(PasswordResetRequest.query.count(), 0)

        admin = User.query.filter_by(role='ADMIN').one()
        with recovery_client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        self.assertEqual(recovery_client.get('/api/admin/password-reset-requests').status_code, 404)
        self.assertEqual(
            recovery_client.post('/api/admin/password-reset-requests/1/issue-code').status_code,
            404,
        )

    def test_database_enforces_one_account_per_role(self):
        db.session.add(User(username='one', full_name='One', email='one@isu.edu.ph', password_hash='x',
                            role='BAC CHAIRMAN', office='CBAO', account_status='PENDING', is_active=False))
        db.session.commit()
        db.session.add(User(username='two', full_name='Two', email='two@isu.edu.ph', password_hash='x',
                            role='BAC CHAIRMAN', office='CBAO', account_status='PENDING', is_active=False))
        with self.assertRaises(IntegrityError):
            db.session.commit()
        db.session.rollback()

    def test_retired_department_endpoints_are_absent_and_normal_modules_remain_accessible(self):
        college = College(code='CED', name='College of Education (CED)')
        aa_office = College(code='AA', name='Academic Affairs')
        db.session.add_all([college, aa_office])
        db.session.flush()
        aa = User(username='academic_affairs', full_name='Academic Affairs', email='aa@isu.edu.ph',
                  password_hash=generate_password_hash('password123'), role='AA', office='AA',
                  college_id=aa_office.id, account_status='ACTIVE', is_active=True)
        college_user = User(username='ced', full_name='CED User', email='ced@isu.edu.ph',
                            password_hash=generate_password_hash('password123'), role='CED',
                            office='CED', college_id=college.id, account_status='ACTIVE', is_active=True)
        master = WishlistItem(office='ADMIN', item_description='Preserved catalog item',
                              category='Office Supplies', unit='piece', quantity=0,
                              estimated_cost=15, status='Available')
        db.session.add_all([aa, college_user, master])
        db.session.commit()
        admin = User.query.filter_by(role='ADMIN').one()
        retired_paths = (
            '/api/department-wishlist',
            '/api/department-wishlist/1/review-copy',
            '/api/department-wishlist/review-copies/1',
        )
        self.assertFalse(any('department-wishlist' in rule.rule for rule in application.app.url_map.iter_rules()))
        for user in (aa, admin, college_user):
            with self.subTest(role=user.role):
                login = self.client.post('/api/auth/login', json={
                    'username': user.username, 'password': 'password123',
                })
                self.assertEqual(login.status_code, 200, login.get_json())
                expected = {'*'} if user.role == 'ADMIN' else set(application.OFFICE_PERMISSIONS)
                self.assertEqual(set(login.get_json()['permissions']), expected)
                self.assertEqual(application.permissions_for(user), expected)
                for path in retired_paths:
                    for method in ('GET', 'POST', 'PATCH'):
                        with self.subTest(path=path, method=method):
                            self.assertEqual(self.client.open(path, method=method, json={}).status_code, 404)
                for path in ('/', '/api/auth/me', '/api/dashboard/ppmp', '/api/ppmp',
                             '/api/ppmp-records', '/api/ppmp-orders', '/api/wishlist',
                             '/api/reports/available', '/api/reports/ppmp-orders',
                             '/api/activity-log', '/api/notifications', '/api/metadata',
                             '/api/auth/available-roles', '/api/colleges'):
                    with self.subTest(path=path):
                        response = self.client.get(path)
                        self.assertEqual(response.status_code, 200, response.get_data(as_text=True)[:500])
                dashboard = self.client.get('/').get_data(as_text=True)
                for removed in ('panel-department-wishlist', 'department-wishlist-nav',
                                '/api/department-wishlist', 'AA Review Copy', 'Mark No PPMP'):
                    self.assertNotIn(removed, dashboard)
                for path in ('/api/apps', '/api/dashboard/app', '/api/admin/pending-users',
                             '/api/admin/activity-log', '/api/admin/audit-log'):
                    expected_status = 200 if user.role == 'ADMIN' else 403
                    with self.subTest(path=path):
                        self.assertEqual(self.client.get(path).status_code, expected_status)
                self.assertEqual(self.client.post('/api/auth/logout').status_code, 200)
        self.assertEqual(db.session.get(WishlistItem, master.id).item_description, 'Preserved catalog item')
        self.assertEqual(PPMPOrder.query.count(), 0)

    def test_retired_department_data_is_preserved_and_not_required_for_startup(self):
        from sqlalchemy import inspect, select
        from models import (
            DepartmentPPMPReviewCopy, DepartmentPPMPReviewItem,
            DepartmentPPMPSubmission, DepartmentPPMPSubmissionItem,
        )

        admin = User.query.filter_by(role='ADMIN').one()
        master = WishlistItem(office='ADMIN', item_description='Preserved catalog item',
                              category='Office Supplies', unit='piece', quantity=0,
                              estimated_cost=15, status='Available')
        db.session.add(master)
        db.session.flush()
        submission = DepartmentPPMPSubmission(
            department_code='CED', fiscal_year=2027, status='Submitted',
            notes='Original historical request', submitted_by=admin.id,
        )
        db.session.add(submission)
        db.session.flush()
        original = DepartmentPPMPSubmissionItem(
            submission_id=submission.id, master_item_id=master.id,
            item_name='Historical item', category='Office Supplies', unit='piece',
            quantity=2, estimated_unit_cost=15, remarks='Original notes',
        )
        review = DepartmentPPMPReviewCopy(
            submission_id=submission.id, status='Reviewed', notes='Historical reviewed copy',
            created_by=admin.id, reviewed_by=admin.id, reviewed_at=datetime(2026, 10, 15, 0, 0),
        )
        db.session.add_all([original, review])
        db.session.flush()
        db.session.add(DepartmentPPMPReviewItem(
            review_copy_id=review.id, master_item_id=master.id, item_name='Historical AA edit',
            category='Office Supplies', unit='piece', quantity=3, estimated_unit_cost=15,
            remarks='Retain edited notes', is_custom_item=False,
        ))
        db.session.commit()
        archive_tables = [model.__table__ for model in (
            DepartmentPPMPSubmission, DepartmentPPMPSubmissionItem,
            DepartmentPPMPReviewCopy, DepartmentPPMPReviewItem,
        )]

        def archive_snapshot():
            return {table.name: [dict(row) for row in db.session.execute(select(table)).mappings()]
                    for table in archive_tables}

        before = archive_snapshot()
        with patch.object(db, 'create_all') as create_all:
            application.initialize_database()
        create_all.assert_not_called()
        self.assertEqual(archive_snapshot(), before)
        self.assertEqual(WishlistItem.query.count(), 1)
        self.assertEqual(User.query.count(), 1)
        # Only this isolated in-memory fixture is dropped to simulate installations
        # that never deployed the retired workflow; production data is never used.
        db.session.remove()
        db.metadata.drop_all(bind=db.engine, tables=list(reversed(archive_tables)))
        without_archive = set(inspect(db.engine).get_table_names())
        with patch.object(db, 'create_all') as create_all:
            application.initialize_database()
        create_all.assert_not_called()
        self.assertEqual(set(inspect(db.engine).get_table_names()), without_archive)
        self.assertTrue({table.name for table in archive_tables}.isdisjoint(without_archive))
        self.assertEqual(WishlistItem.query.count(), 1)
        self.assertEqual(User.query.count(), 1)
        application.initialize_database(apply_migrations=True)
        self.assertTrue({table.name for table in archive_tables}.isdisjoint(
            set(inspect(db.engine).get_table_names())))
        self.assertEqual(WishlistItem.query.count(), 1)
        self.assertEqual(User.query.count(), 1)

    def test_office_user_cannot_create_for_another_college(self):
        first = College(code='A', name='College A')
        second = College(code='B', name='College B')
        db.session.add_all([first, second]); db.session.flush()
        user = User(username='college_a', full_name='A User', email='a@isu.edu.ph', password_hash='x',
                    role='CBAO', office='CBAO', college_id=first.id, account_status='ACTIVE', is_active=True)
        db.session.add(user); db.session.commit()
        with self.client.session_transaction() as session:
            session['user_id'] = user.id
            session['role'] = user.role
        response = self.client.post('/api/ppmp', json={'college_id': second.id, 'item_description': 'Paper',
            'unit': 'ream', 'quantity': 1, 'unit_cost': 100, 'fund_source': 'GAA',
            'mode_of_procurement': 'Public Bidding', 'procurement_type': 'Goods'})
        self.assertEqual(response.status_code, 403)

    def test_ppmp_order_uses_master_and_custom_items_without_polluting_wishlist(self):
        office = College(code='OFFICE-A', name='Office A')
        db.session.add(office)
        db.session.flush()
        office_user = User(username='office_a', full_name='Office A User', email='office-a@isu.edu.ph',
                           password_hash='x', role='OFFICE-A', office='OFFICE-A', college_id=office.id,
                           account_status='ACTIVE', is_active=True)
        eo = User(username='eo', full_name='EO', email='eo@isu.edu.ph', password_hash='x', role='EO',
                  office='EO', account_status='ACTIVE', is_active=True)
        procurement = User(username='procurement_flow', full_name='Procurement', email='proc-flow@isu.edu.ph',
                           password_hash='x', role='PROCUREMENT', office='PROCUREMENT',
                           account_status='ACTIVE', is_active=True)
        db.session.add_all([office_user, eo, procurement])
        db.session.commit()

        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        master_response = self.client.post('/api/wishlist', json={
            'item_description': 'Printer paper', 'category': 'Office Supplies',
            'unit': 'ream', 'estimated_cost': 250,
        })
        self.assertEqual(master_response.status_code, 201)
        master_id = master_response.get_json()['id']
        updated_master = self.client.put(f'/api/wishlist/{master_id}', json={
            'item_description': 'Printer paper A4',
        })
        self.assertEqual(updated_master.status_code, 200)
        self.assertEqual(updated_master.get_json()['item_description'], 'Printer paper A4')

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = office_user.id, office_user.role
        forbidden_master = self.client.post('/api/wishlist', json={
            'item_description': 'Office-created catalog item', 'category': 'Other',
            'unit': 'piece', 'estimated_cost': 10,
        })
        self.assertEqual(forbidden_master.status_code, 403)
        self.assertEqual(self.client.put(
            f'/api/wishlist/{master_id}', json={'item_description': 'Unauthorized edit'},
        ).status_code, 403)
        self.assertEqual(self.client.delete(f'/api/wishlist/{master_id}').status_code, 403)
        created = self.client.post('/api/ppmp-orders', json={
            'fiscal_year': 2026, 'quarter': 'Q1 (Jan-Mar)',
            'items': [
                {'master_item_id': master_id, 'quantity': 5},
                {
                    'item_name': 'Custom archival box', 'description': 'Acid-free box',
                    'category': 'Goods', 'unit': 'box', 'quantity': 3,
                    'estimated_unit_cost': 400,
                },
            ],
        })
        self.assertEqual(created.status_code, 201)
        order = created.get_json()
        order_id = order['id']
        self.assertEqual(order['status'], 'Submitted')
        self.assertEqual(order['current_stage'], 'APP / Procurement Consolidation')
        self.assertEqual(len(order['items']), 2)
        self.assertIsNone(order['fund_source'])
        self.assertIsNone(order['mode_of_procurement'])
        self.assertIsNone(order['procurement_type'])
        self.assertIsNone(order['remarks'])
        self.assertEqual(WishlistItem.query.count(), 1)
        self.assertEqual(PPMPItem.query.count(), 0)
        stored_custom = PPMPOrderItem.query.filter_by(order_id=order_id, is_custom_item=True).one()
        self.assertIsNone(stored_custom.master_item_id)
        self.assertEqual(PPMPNotification.query.filter_by(
            recipient_id=procurement.id, title='Wishlist request ready for APP consolidation',
        ).count(), 1)
        for index, item in enumerate(order['items']):
            report_update = self.client.put(
                f'/api/reports/ppmp-orders/{order_id}/items/{item["id"]}',
                json={
                    'project_type': 'Goods/Services',
                    'procurement_method': application.PPMP_REPORT_PROCUREMENT_METHODS[index],
                    'pre_procurement_conference': 'No',
                    'procurement_start': '2026-01',
                    'procurement_end': '2026-03',
                    'delivery_period': '2026-04',
                    'source_of_funds': 'GAA/101',
                    'remarks': '' if index else 'Quarterly office requirements',
                },
            )
            self.assertEqual(report_update.status_code, 200)

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = procurement.id, procurement.role
        queue = self.client.get('/api/ppmp-orders?status=Submitted')
        self.assertEqual(queue.status_code, 200)
        self.assertEqual([row['id'] for row in queue.get_json()], [order_id])
        missing_reason = self.client.post(f'/api/ppmp-orders/{order_id}/transition', json={'action': 'return'})
        self.assertEqual(missing_reason.status_code, 400)
        consolidated = self.client.post(f'/api/ppmp-orders/{order_id}/transition', json={'action': 'approve'})
        self.assertEqual(consolidated.status_code, 200)
        self.assertEqual(consolidated.get_json()['status'], 'Pending Admin Approval')
        self.assertEqual(PPMPItem.query.count(), 0)

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        approved = self.client.post(f'/api/ppmp-orders/{order_id}/transition', json={'action': 'approve'})
        self.assertEqual(approved.status_code, 200)
        self.assertEqual(approved.get_json()['status'], 'Approved')
        self.assertEqual(WishlistItem.query.count(), 1)
        self.assertEqual(PPMPItem.query.count(), 2)
        self.assertEqual({item.created_by for item in PPMPItem.query.all()}, {office_user.id})
        self.assertEqual({item.approved_by for item in PPMPItem.query.all()}, {admin.id})
        self.assertTrue(all(item.ppmp_item_id for item in PPMPOrderItem.query.filter_by(order_id=order_id)))
        self.assertEqual(PPMPNotification.query.filter_by(
            recipient_id=office_user.id, title='Wishlist request approved',
        ).count(), 1)
        linked_delete = self.client.delete(f'/api/wishlist/{master_id}')
        self.assertEqual(linked_delete.status_code, 409)
        self.assertEqual(WishlistItem.query.count(), 1)

        report = self.client.get('/api/reports/ppmp-orders')
        self.assertEqual(report.status_code, 200)
        self.assertEqual(report.get_json()['summary']['total_requests'], 1)
        self.assertEqual(report.get_json()['summary']['total_ordered_items'], 8)
        self.assertEqual(report.get_json()['source_distribution']['CSE Master List'], 1)
        self.assertEqual(report.get_json()['source_distribution']['Custom / Unlisted'], 1)

    def test_cse_options_reports_and_office_scoping(self):
        first = College(code='REPORT-A', name='Report Office A')
        second = College(code='REPORT-B', name='Report Office B')
        db.session.add_all([first, second])
        db.session.flush()
        first_user = User(
            username='report_a', full_name='Report A User', email='report-a@isu.edu.ph',
            password_hash='x', role='REPORT-A', office='REPORT-A', college_id=first.id,
            account_status='ACTIVE', is_active=True,
        )
        second_user = User(
            username='report_b', full_name='Report B User', email='report-b@isu.edu.ph',
            password_hash='x', role='REPORT-B', office='REPORT-B', college_id=second.id,
            account_status='ACTIVE', is_active=True,
        )
        db.session.add_all([first_user, second_user])
        admin = User.query.filter_by(role='ADMIN').one()
        master = WishlistItem(
            office='ADMINISTRATION', created_by=admin.id, item_description='CSE paper',
            category='Goods', unit='ream', quantity=25, estimated_cost=250, status='Available',
        )
        db.session.add(master)
        db.session.flush()
        first_order = PPMPOrder(
            reference_number='REPORT-A-2026', office_id=first.id, requested_by=first_user.id,
            fiscal_year=2026, fund_source='GAA/101', mode_of_procurement='Public Bidding',
            procurement_type='Goods/Services', quarter='Q1 (Jan-Mar)', status='Submitted',
        )
        second_order = PPMPOrder(
            reference_number='REPORT-B-2026', office_id=second.id, requested_by=second_user.id,
            fiscal_year=2026, fund_source='161', mode_of_procurement='Public Bidding',
            procurement_type='Infrastructure', quarter='Q1 (Jan-Mar)', status='Submitted',
        )
        db.session.add_all([first_order, second_order])
        db.session.flush()
        db.session.add_all([
            PPMPOrderItem(
                order_id=first_order.id, master_item_id=master.id, item_name='CSE paper',
                category='Goods', quantity=3, unit='ream', estimated_unit_cost=250,
                estimated_budget=750, is_custom_item=False,
            ),
            PPMPOrderItem(
                order_id=second_order.id, item_name='Other infrastructure item',
                category='Infrastructure', quantity=1, unit='lot', estimated_unit_cost=1000,
                estimated_budget=1000, is_custom_item=True,
            ),
        ])
        db.session.commit()

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        metadata = self.client.get('/api/metadata').get_json()
        self.assertEqual(metadata['ppmp_order_categories'], ['Goods', 'Infrastructure', 'Consulting Services'])
        self.assertEqual(metadata['ppmp_order_fund_sources'], ['GAA/101', 'SB 164', '161', '601'])
        self.assertEqual(metadata['ppmp_order_types'], ['Goods/Services', 'Infrastructure', 'Consulting Services'])

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        publication = self.client.post('/api/reports/available', json={
            'title': 'Office A Updated PPMP Report', 'office_id': first.id, 'fiscal_year': 2026,
        })
        self.assertEqual(publication.status_code, 201)
        report_id = publication.get_json()['id']
        updated = self.client.put(f'/api/reports/available/{report_id}', json={
            'title': 'Office A Updated PPMP Report', 'office_id': first.id, 'fiscal_year': 2026,
        })
        self.assertEqual(updated.status_code, 200)

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = first_user.id, first_user.role
        available = self.client.get('/api/reports/available')
        self.assertEqual([row['id'] for row in available.get_json()], [report_id])
        scoped = self.client.get(f'/api/reports/ppmp-orders?report_id={report_id}')
        self.assertEqual(scoped.status_code, 200)
        self.assertEqual([row['reference_number'] for row in scoped.get_json()['records']], ['REPORT-A-2026'])
        self.assertEqual(self.client.get(
            f'/api/reports/ppmp-orders?office_id={second.id}'
        ).status_code, 403)
        pdf = self.client.get(f'/api/reports/ppmp-orders/export/pdf?report_id={report_id}')
        excel = self.client.get(f'/api/reports/ppmp-orders/export/excel?report_id={report_id}')
        self.assertEqual((pdf.status_code, pdf.mimetype), (200, 'application/pdf'))
        self.assertEqual(
            (excel.status_code, excel.mimetype),
            (200, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
        )

        invalid_base = {
            'fiscal_year': 2026, 'quarter': 'Q1 (Jan-Mar)',
            'items': [{'item_name': 'Other item', 'category': 'Invalid', 'unit': 'piece',
                       'quantity': 1, 'estimated_unit_cost': 1}],
        }
        self.assertEqual(self.client.post('/api/ppmp-orders', json=invalid_base).status_code, 400)
        missing_year = dict(invalid_base)
        missing_year.pop('fiscal_year')
        missing_year['items'] = [dict(invalid_base['items'][0], category='Goods')]
        self.assertEqual(self.client.post('/api/ppmp-orders', json=missing_year).status_code, 400)
        invalid_quarter = dict(invalid_base, quarter='Q5')
        invalid_quarter['items'] = [dict(invalid_base['items'][0], category='Goods')]
        self.assertEqual(self.client.post('/api/ppmp-orders', json=invalid_quarter).status_code, 400)

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = second_user.id, second_user.role
        self.assertEqual(self.client.get(f'/api/reports/ppmp-orders?report_id={report_id}').status_code, 404)
        self.assertEqual(self.client.get('/api/reports/available').get_json(), [])

    def test_user_opens_one_pending_request_and_exports_exact_12_column_report(self):
        first = College(code='REPORT-USER', name='Budget Office')
        second = College(code='REPORT-OTHER', name='Procurement Office')
        db.session.add_all([first, second])
        db.session.flush()
        first_user = User(
            username='report_user', full_name='Report User', email='report-user@isu.edu.ph',
            password_hash='x', role='REPORT-USER', office='REPORT-USER', college_id=first.id,
            account_status='ACTIVE', is_active=True,
        )
        second_user = User(
            username='report_other', full_name='Other User', email='report-other@isu.edu.ph',
            password_hash='x', role='REPORT-OTHER', office='REPORT-OTHER', college_id=second.id,
            account_status='ACTIVE', is_active=True,
        )
        db.session.add_all([first_user, second_user])
        db.session.flush()
        selected = PPMPOrder(
            reference_number='PENDING-REPORT-2027', office_id=first.id, requested_by=first_user.id,
            fiscal_year=2027, fund_source='GAA/101', mode_of_procurement='Public Bidding',
            procurement_type='Goods/Services', quarter='Q1 (Jan-Mar)', status='Submitted',
        )
        other = PPMPOrder(
            reference_number='OTHER-REPORT-2027', office_id=second.id, requested_by=second_user.id,
            fiscal_year=2027, fund_source='161', mode_of_procurement='Public Bidding',
            procurement_type='Goods/Services', quarter='Q1 (Jan-Mar)', status='Submitted',
        )
        db.session.add_all([selected, other])
        db.session.flush()
        selected_items = [
            PPMPOrderItem(
                order_id=selected.id, item_name='Requested Item 1', category='Goods',
                quantity=10, unit='pcs', estimated_unit_cost=500, estimated_budget=5000,
                supporting_documents='Technical specifications',
                created_by=first_user.id, is_custom_item=True,
            ),
            PPMPOrderItem(
                order_id=selected.id, item_name='Requested Item 2', category='Infrastructure',
                quantity=2, unit='sets', estimated_unit_cost=2500, estimated_budget=5000,
                supporting_documents='Program of works',
                created_by=first_user.id, is_custom_item=True,
            ),
        ]
        other_item = PPMPOrderItem(
            order_id=other.id, item_name='Other Office Private Item', category='Goods',
            quantity=1, unit='unit', estimated_unit_cost=1, estimated_budget=1,
            created_by=second_user.id, is_custom_item=True,
        )
        db.session.add_all([*selected_items, other_item])
        db.session.commit()

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = first_user.id, first_user.role
        listing = self.client.get('/api/reports/ppmp-orders').get_json()
        self.assertEqual([row['reference_number'] for row in listing['records']], ['PENDING-REPORT-2027'])
        self.assertTrue(listing['records'][0]['is_pending_report'])
        detail = self.client.get(f'/api/reports/ppmp-orders/{selected.id}')
        self.assertEqual(detail.status_code, 200)
        detail_json = detail.get_json()
        self.assertEqual(detail_json['fiscal_year'], 2027)
        self.assertEqual(detail_json['end_user_or_implementing_unit'], 'Budget Office')
        self.assertTrue(detail_json['editable'])
        self.assertEqual(
            [item['general_description'] for item in detail_json['items']],
            ['Requested Item 1', 'Requested Item 2'],
        )
        self.assertNotIn('Other Office Private Item', str(detail_json))
        self.assertEqual(tuple(detail_json['procurement_methods']), application.PPMP_REPORT_PROCUREMENT_METHODS)
        self.assertEqual(len(detail_json['procurement_methods']), 26)
        self.assertEqual(detail_json['procurement_types'], ['Goods/Services', 'Infrastructure', 'Consulting Services'])
        self.assertEqual(detail_json['fund_sources'], ['GAA/101', 'SB 164', '161', '601'])
        self.assertEqual(
            self.client.get(f'/api/reports/ppmp-orders/{other.id}').status_code, 403,
        )
        self.assertEqual(
            self.client.put(
                f'/api/reports/ppmp-orders/{selected.id}/items/{selected_items[0].id}',
                json={'procurement_method': 'Public Bidding'},
            ).status_code,
            400,
        )
        self.assertEqual(
            self.client.put(
                f'/api/reports/ppmp-orders/{selected.id}/items/{selected_items[0].id}',
                json={'estimated_budget': 1},
            ).status_code,
            400,
        )
        self.assertEqual(self.client.put(
            f'/api/reports/ppmp-orders/{selected.id}/items/{selected_items[0].id}',
            json={'project_type': 'Goods'},
        ).status_code, 400)
        self.assertEqual(self.client.put(
            f'/api/reports/ppmp-orders/{selected.id}/items/{selected_items[0].id}',
            json={'source_of_funds': 'GAA'},
        ).status_code, 400)
        self.assertEqual(
            self.client.get(f'/api/reports/ppmp-orders/{selected.id}/export/pdf').status_code,
            409,
        )
        for index, item in enumerate(selected_items):
            response = self.client.put(
                f'/api/reports/ppmp-orders/{selected.id}/items/{item.id}',
                json={
                    'project_type': application.PPMP_ORDER_TYPES[index],
                    'procurement_method': application.PPMP_REPORT_PROCUREMENT_METHODS[index],
                    'pre_procurement_conference': 'Yes' if index == 0 else 'No',
                    'procurement_start': '2026-02',
                    'procurement_end': '2026-12',
                    'delivery_period': '2027-01',
                    'source_of_funds': application.PPMP_ORDER_FUND_SOURCES[index],
                    'remarks': 'Updated item remark' if index == 0 else '',
                },
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.get_json()['remarks'], 'Updated item remark' if index == 0 else '')

        pdf = self.client.get(f'/api/reports/ppmp-orders/{selected.id}/export/pdf')
        excel = self.client.get(f'/api/reports/ppmp-orders/{selected.id}/export/excel')
        self.assertEqual((pdf.status_code, pdf.mimetype), (200, 'application/pdf'))
        self.assertEqual(
            (excel.status_code, excel.mimetype),
            (200, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
        )
        pdf_text = '\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(pdf.data)).pages)
        normalized_pdf = re.sub(r'\s+', ' ', pdf_text)
        self.assertIn('PROJECT PROCUREMENT MANAGEMENT PLAN (PPMP) NO. 1', normalized_pdf)
        self.assertIn('Fiscal Year: 2027', normalized_pdf)
        self.assertIn('End-User or Implementing Unit: Budget Office', normalized_pdf)
        self.assertIn('Requested Item 1', normalized_pdf)
        self.assertIn('Requested Item 2', normalized_pdf)
        self.assertNotIn('Other Office Private Item', normalized_pdf)
        self.assertIn(b'/Subtype /Image', pdf.data)

        sheet = load_workbook(io.BytesIO(excel.data), data_only=True).active
        self.assertEqual(sheet.max_column, 12)
        self.assertEqual(sheet['A1'].value, 'Fiscal Year:')
        self.assertEqual(sheet['B1'].value, 2027)
        self.assertEqual(sheet['A2'].value, 'End-User or Implementing Unit:')
        self.assertEqual(sheet['B2'].value, 'Budget Office')
        self.assertEqual(tuple(cell.value for cell in sheet[4]), application.PPMP_REPORT_COLUMNS)
        self.assertEqual(sheet['A5'].value, 'Requested Item 1')
        self.assertEqual(sheet['A6'].value, 'Requested Item 2')
        self.assertEqual(sheet['B5'].value, 'Goods/Services')
        self.assertEqual(sheet['B6'].value, 'Infrastructure')
        self.assertEqual(sheet['I5'].value, 'GAA/101')
        self.assertEqual(sheet['I6'].value, 'SB 164')
        self.assertEqual(sheet['J5'].value, 5000)
        self.assertEqual(sheet['L5'].value, 'Updated item remark')
        self.assertIsNone(sheet['L6'].value)
        self.assertNotIn('Other Office Private Item', str(list(sheet.values)))

    def test_approved_12_column_wishlist_item_creates_a_separate_admin_request(self):
        office = College(code='SUPPLY', name='Supply Office')
        db.session.add(office)
        db.session.flush()
        office_user = User(
            username='supply_user', full_name='Supply User', email='supply@isu.edu.ph',
            password_hash='x', role='SUPPLY', office='SUPPLY', college_id=office.id,
            account_status='ACTIVE', is_active=True,
        )
        eo = User(
            username='eo_review', full_name='EO Review', email='eo-review@isu.edu.ph',
            password_hash='x', role='EO', office='EO', account_status='ACTIVE', is_active=True,
        )
        procurement = User(
            username='proc_review', full_name='Procurement Review', email='proc-review@isu.edu.ph',
            password_hash='x', role='PROCUREMENT', office='PROCUREMENT',
            account_status='ACTIVE', is_active=True,
        )
        db.session.add_all([office_user, eo, procurement])
        db.session.commit()

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = office_user.id, office_user.role
        created_record = self.client.post('/api/ppmp-records', json={'fiscal_year': 2026})
        self.assertEqual(created_record.status_code, 201)
        record_id = created_record.get_json()['id']
        row_payload = {
            'general_description': 'Office supplies',
            'project_type': 'Goods',
            'quantity': 3,
            'unit_size': 'lots',
            'procurement_method': 'Small Value Procurement',
            'pre_procurement_conference': 'No',
            'procurement_start': '2026-01',
            'procurement_end': '2026-01',
            'delivery_period': '2026-01',
            'schedule_as_needed': False,
            'source_of_funds': '164',
            'estimated_budget': '20000.00',
            'supporting_documents': 'CSE to be procured at PS-DBM',
            'remarks': '',
        }
        added = self.client.post(f'/api/ppmp-records/{record_id}/items', json=row_payload)
        self.assertEqual(added.status_code, 201)
        wishlist_item_id = added.get_json()['id']
        self.assertEqual(PPMPOrder.query.count(), 0, 'Adding a Wishlist row must not create an order.')
        submitted = self.client.post(f'/api/ppmp-records/{record_id}/submit')
        self.assertEqual(submitted.get_json()['status'], 'Submitted')

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = eo.id, eo.role
        self.assertEqual(
            self.client.post(f'/api/ppmp-records/{record_id}/transition', json={'action': 'approve'}).get_json()['status'],
            'Under Review',
        )
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = procurement.id, procurement.role
        self.assertEqual(
            self.client.post(f'/api/ppmp-records/{record_id}/transition', json={'action': 'approve'}).get_json()['status'],
            'Pending Admin Approval',
        )
        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        approved_record = self.client.post(
            f'/api/ppmp-records/{record_id}/transition', json={'action': 'approve'},
        )
        self.assertEqual(approved_record.get_json()['status'], 'Approved')
        self.assertEqual(PPMPItem.query.count(), 1)

        other_office = College(code='RECORDS', name='Records Office')
        other_user = User(
            username='records_user', full_name='Records User', email='records@isu.edu.ph',
            password_hash='x', role='RECORDS', office='RECORDS', college=other_office,
            account_status='ACTIVE', is_active=True,
        )
        db.session.add_all([other_office, other_user])
        db.session.commit()
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = other_user.id, other_user.role
        forbidden = self.client.post('/api/ppmp-orders', json={
            'wishlist_item_ids': [wishlist_item_id],
        })
        self.assertEqual(forbidden.status_code, 403)

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = office_user.id, office_user.role
        order_response = self.client.post('/api/ppmp-orders', json={
            'wishlist_item_ids': [wishlist_item_id],
            'remarks': 'Request approved PPMP item',
        })
        self.assertEqual(order_response.status_code, 201)
        order = order_response.get_json()
        self.assertEqual(order['status'], 'Submitted')
        self.assertEqual(order['items'][0]['source_type'], 'PPMP Wishlist Item')
        self.assertEqual(order['items'][0]['wishlist_item_id'], wishlist_item_id)
        self.assertEqual(PPMPItem.query.count(), 1)
        self.assertEqual(
            self.client.post('/api/ppmp-orders', json={'wishlist_item_ids': [wishlist_item_id]}).status_code,
            409,
        )

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = procurement.id, procurement.role
        consolidated_order = self.client.post(
            f"/api/ppmp-orders/{order['id']}/transition", json={'action': 'approve'},
        )
        self.assertEqual(consolidated_order.status_code, 200)
        self.assertEqual(consolidated_order.get_json()['status'], 'Pending Admin Approval')

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        approved_order = self.client.post(
            f"/api/ppmp-orders/{order['id']}/transition", json={'action': 'approve'},
        )
        self.assertEqual(approved_order.status_code, 200)
        self.assertEqual(approved_order.get_json()['status'], 'Approved')
        self.assertEqual(PPMPItem.query.count(), 1, 'Order approval must not duplicate the PPMP item.')

    def test_app_creation_without_batch_returns_validation_error(self):
        user = User(username='procurement', full_name='Procurement', email='procurement@isu.edu.ph',
                    password_hash='x', role='PROCUREMENT', office='PROCUREMENT',
                    account_status='ACTIVE', is_active=True)
        db.session.add(user)
        db.session.commit()
        with self.client.session_transaction() as session:
            session['user_id'] = user.id
            session['role'] = user.role
        response = self.client.post('/api/apps', json={
            'office': 'PROCUREMENT', 'fiscal_year': 2026, 'plan_type': 'APP'
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn('consolidation batch', response.get_json()['error'].lower())

    def test_normal_startup_never_runs_schema_creation(self):
        connection = MagicMock()
        with patch.object(application, 'get_db', return_value=connection), \
             patch.object(application, 'verify_database_schema') as verify, \
             patch.object(db, 'create_all') as create_all:
            application.initialize_database()
        connection.close.assert_called_once_with()
        verify.assert_called_once_with()
        create_all.assert_not_called()

    def test_office_management_routes_are_removed_but_directory_remains_read_only(self):
        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'] = admin.id
            session['role'] = admin.role
        self.assertEqual(self.client.get('/api/admin/offices').status_code, 404)
        self.assertEqual(self.client.post('/api/admin/offices', json={}).status_code, 404)
        self.assertEqual(self.client.put('/api/admin/offices/1', json={}).status_code, 404)
        self.assertEqual(self.client.get('/api/colleges').status_code, 200)

    def test_account_approval_links_a_configured_office_without_management_ui(self):
        registered = self.registration(
            email='cbao@isu.edu.ph', username='cbao', role='CBAO', office='CBAO',
        )
        self.assertEqual(registered.status_code, 201)
        pending = User.query.filter_by(role='CBAO').one()
        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        approved = self.client.post(f'/api/admin/pending-users/{pending.id}/approve')
        self.assertEqual(approved.status_code, 200)
        office = College.query.filter_by(code='CBAO').one()
        self.assertEqual(approved.get_json()['user']['college_id'], office.id)
        self.assertTrue(office.is_active)

    def test_registration_status_flow_uses_signed_reference_and_system_notification(self):
        login_html = self.client.get('/').get_data(as_text=True)
        self.assertNotIn('<label for="role">Available role</label>', login_html)
        self.assertIn('<label for="role">Role</label>', login_html)
        self.assertIn('id="register-confirm"', login_html)
        self.assertIn('Check Account Status', login_html)
        self.assertIn('Registration Successful', login_html)

        registered = self.registration()
        self.assertEqual(registered.status_code, 201)
        registration_token = registered.get_json()['registration_token']
        self.assertTrue(registration_token)
        pending = self.client.post(
            '/api/auth/registration-status', json={'registration_token': registration_token},
        )
        self.assertEqual(pending.status_code, 200)
        self.assertEqual(pending.get_json()['account_status'], 'PENDING')
        self.assertNotIn('user_id', pending.get_json())

        blocked = self.client.post('/api/auth/login', json={
            'username': 'procurement', 'password': 'password123',
        })
        self.assertEqual(blocked.status_code, 403)
        self.assertEqual(blocked.get_json()['code'], 'ACCOUNT_PENDING')

        pending_user = User.query.filter_by(username='procurement').one()
        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        pending_list = self.client.get('/api/admin/pending-users').get_json()['users'][0]
        self.assertIn('created_at', pending_list)
        self.assertEqual(self.client.post(
            f'/api/admin/pending-users/{pending_user.id}/approve',
        ).status_code, 200)
        self.assertEqual(PPMPNotification.query.filter_by(
            recipient_id=pending_user.id, title='Account approved',
        ).count(), 1)

        self.client.post('/api/auth/logout')
        active = self.client.post(
            '/api/auth/registration-status', json={'registration_token': registration_token},
        )
        self.assertEqual(active.get_json()['account_status'], 'ACTIVE')
        signed_in = self.client.post('/api/auth/login', json={
            'username': 'procurement', 'password': 'password123',
        })
        self.assertEqual(signed_in.status_code, 200)
        self.assertIn('app.create', signed_in.get_json()['permissions'])

    def test_profile_activity_log_records_successful_sessions_per_user(self):
        admin = User.query.filter_by(username='admin').one()
        other_user = User(
            username='activity_user', full_name='Activity User', email='activity@isu.edu.ph',
            password_hash=generate_password_hash('password123'), role='EO', office='EO',
            account_status='ACTIVE', is_active=True,
        )
        db.session.add(other_user)
        db.session.commit()

        failed = self.client.post('/api/auth/login', json={
            'username': admin.username, 'password': 'incorrect-password',
        })
        self.assertEqual(failed.status_code, 401)
        self.assertEqual(UserActivityLog.query.count(), 0)

        signed_in = self.client.post('/api/auth/login', json={
            'username': admin.username, 'password': 'password123',
        })
        self.assertEqual(signed_in.status_code, 200)
        row = UserActivityLog.query.one()
        self.assertEqual(row.user_id, admin.id)
        self.assertIsNone(row.time_out)
        profile = signed_in.get_json()['user']
        self.assertEqual(profile['office'], 'ADMINISTRATION')
        self.assertEqual(profile['full_name'], 'Admin')
        self.assertEqual(profile['username'], 'admin')
        self.assertEqual(profile['email'], 'admin@test.local')

        history = self.client.get('/api/activity-log')
        self.assertEqual(history.status_code, 200)
        self.assertEqual(set(history.get_json()[0]), {'date', 'time_in', 'time_out'})
        self.assertEqual(history.get_json()[0]['time_out'], '')
        self.assertEqual(self.client.post('/api/activity-log').status_code, 405)
        self.assertEqual(self.client.put('/api/activity-log').status_code, 405)
        self.assertEqual(self.client.delete('/api/activity-log').status_code, 405)

        dashboard = self.client.get('/').get_data(as_text=True)
        self.assertIn('id="profile-panel"', dashboard)
        self.assertIn('id="activity-log-nav"', dashboard)
        self.assertIn('<th>Date</th><th>Time In</th><th>Time Out</th>', dashboard)
        for field_id in ('profile-office', 'profile-full-name', 'profile-username', 'profile-gmail'):
            self.assertIn(f'id="{field_id}"', dashboard)

        self.assertEqual(self.client.post('/api/auth/logout').status_code, 200)
        db.session.refresh(row)
        self.assertIsNotNone(row.time_out)

        other_client = application.app.test_client()
        other_login = other_client.post('/api/auth/login', json={
            'username': other_user.username, 'password': 'password123',
        })
        self.assertEqual(other_login.status_code, 200)
        own_history = other_client.get('/api/activity-log').get_json()
        self.assertEqual(len(own_history), 1)
        self.assertEqual(UserActivityLog.query.count(), 2)
        self.assertEqual(UserActivityLog.query.filter_by(user_id=admin.id).count(), 1)
        self.assertEqual(UserActivityLog.query.filter_by(user_id=other_user.id).count(), 1)

    def test_admin_activity_tracker_filters_office_and_denies_non_admin(self):
        cbao_user = User(
            username='cbao_user', full_name='CBAO User', email='cbao@isu.edu.ph',
            password_hash='x', role='CBAO', office='CBAO',
            account_status='ACTIVE', is_active=True,
        )
        cbao_reviewer = User(
            username='cbao_reviewer', full_name='CBAO Reviewer', email='cbao-reviewer@isu.edu.ph',
            password_hash='x', role='EO', office='CBAO',
            account_status='ACTIVE', is_active=True,
        )
        budget_user = User(
            username='budget_user', full_name='Budget User', email='budget@isu.edu.ph',
            password_hash='x', role='BUDGET', office='BUDGET',
            account_status='ACTIVE', is_active=True,
        )
        db.session.add_all([cbao_user, cbao_reviewer, budget_user])
        db.session.flush()
        db.session.add_all([
            UserActivityLog(user_id=cbao_user.id, time_in=datetime(2026, 1, 15, 1, 30),
                            time_out=datetime(2026, 1, 15, 3, 45)),
            UserActivityLog(user_id=cbao_reviewer.id, time_in=datetime(2026, 2, 2, 4, 5)),
            UserActivityLog(user_id=budget_user.id, time_in=datetime(2026, 3, 3, 6, 7)),
        ])
        db.session.commit()

        with self.client.session_transaction() as session:
            admin = User.query.filter_by(role='ADMIN').one()
            session['user_id'], session['role'] = admin.id, admin.role
        office_list = self.client.get('/api/admin/activity-log').get_json()
        self.assertIn('CBAO', office_list['offices'])
        self.assertEqual(office_list['records'], [])

        cbao_records = self.client.get('/api/admin/activity-log?office=CBAO').get_json()['records']
        self.assertEqual(len(cbao_records), 2)
        self.assertEqual(cbao_records[0]['user'], 'CBAO Reviewer')
        self.assertEqual(cbao_records[0]['login_month'], 'February')
        self.assertEqual(cbao_records[0]['login_day'], '02')
        self.assertEqual(cbao_records[0]['login_year'], '2026')
        self.assertEqual(cbao_records[0]['login_time'], '12:05:00 PM')
        self.assertEqual(cbao_records[0]['logout_month'], '')
        self.assertEqual(cbao_records[1]['user'], 'CBAO User')
        self.assertEqual(cbao_records[1]['login_month'], 'January')
        self.assertEqual(cbao_records[1]['login_day'], '15')
        self.assertEqual(cbao_records[1]['login_year'], '2026')
        self.assertEqual(cbao_records[1]['login_time'], '09:30:00 AM')
        self.assertEqual(cbao_records[1]['logout_month'], 'January')
        self.assertEqual(cbao_records[1]['logout_day'], '15')
        self.assertEqual(cbao_records[1]['logout_year'], '2026')
        self.assertEqual(cbao_records[1]['logout_time'], '11:45:00 AM')

        user_client = application.app.test_client()
        with user_client.session_transaction() as session:
            session['user_id'], session['role'] = cbao_user.id, cbao_user.role
        self.assertEqual(user_client.get('/api/admin/activity-log').status_code, 403)
        self.assertEqual(user_client.get('/api/admin/activity-log?office=BUDGET').status_code, 403)
        user_html = user_client.get('/').get_data(as_text=True)
        self.assertNotIn('Activity Log Tracker', user_html)
        admin_html = self.client.get('/').get_data(as_text=True)
        self.assertIn('Activity Log Tracker', admin_html)
        self.assertIn('<th class="activity-tracker-title" colspan="5">Log In</th>', admin_html)
        self.assertIn('<th>User</th><th>Login Month</th><th>Login Day</th><th>Login Year</th><th>Login Time</th>', admin_html)
        self.assertIn('<th class="activity-tracker-title" colspan="4">Log Out</th>', admin_html)
        self.assertIn('<th>Logout Month</th><th>Logout Day</th><th>Logout Year</th><th>Logout Time</th>', admin_html)
        self.assertIn('<th class="activity-tracker-title" colspan="2">Other Activity Conducted</th>', admin_html)

    def test_app_12_column_sections_totals_workflow_permissions_and_exports(self):
        office = College(code='APP-OFFICE', name='APP Implementing Office')
        procurement = User(
            username='app_procurement', full_name='APP Procurement', email='app-procurement@isu.edu.ph',
            password_hash='x', role='PROCUREMENT', office='PROCUREMENT',
            account_status='ACTIVE', is_active=True,
        )
        bac = User(username='app_bac', full_name='BAC Reviewer', email='bac@isu.edu.ph', password_hash='x',
                   role='BAC CHAIRMAN', office='BAC CHAIRMAN', account_status='ACTIVE', is_active=True)
        eo = User(username='app_eo', full_name='EO Reviewer', email='app-eo@isu.edu.ph', password_hash='x',
                  role='EO', office='EO', account_status='ACTIVE', is_active=True)
        budget = User(username='app_budget', full_name='Budget Reviewer', email='app-budget@isu.edu.ph',
                      password_hash='x', role='BUDGET', office='BUDGET', account_status='ACTIVE', is_active=True)
        db.session.add_all([office, procurement, bac, eo, budget])
        db.session.flush()
        ppmp_item = PPMPItem(
            college_id=office.id, fiscal_year=2026, item_description='Legitimate approved source project',
            unit='lot', quantity=1, unit_cost=1250, authorized_budget=1250,
            fund_source='GAA', mode_of_procurement='Small Value Procurement',
            procurement_type='Goods', status='Approved', created_by=procurement.id,
        )
        batch = ConsolidationBatch(fiscal_year=2026, app_cycle='ANNUAL', created_by=procurement.id)
        db.session.add_all([ppmp_item, batch]); db.session.flush()
        db.session.add(ConsolidationSource(batch_id=batch.id, ppmp_item_id=ppmp_item.id))
        db.session.commit()

        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        admin_create = self.client.post('/api/apps', json={
            'office': office.code, 'fiscal_year': 2026, 'plan_type': 'APP',
            'consolidation_batch_id': batch.id,
        })
        self.assertEqual(admin_create.status_code, 403)

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = procurement.id, procurement.role
        created = self.client.post('/api/apps', json={
            'office': office.code, 'fiscal_year': 2026, 'plan_type': 'APP',
            'consolidation_batch_id': batch.id,
        })
        self.assertEqual(created.status_code, 201)
        plan_id = created.get_json()['id']
        self.assertEqual(created.get_json()['item_count'], 1)
        shared_records = self.client.get('/api/apps')
        self.assertEqual(shared_records.status_code, 200)
        self.assertEqual([row['id'] for row in shared_records.get_json()], [plan_id])
        app_dashboard = self.client.get('/api/dashboard/app')
        self.assertEqual(app_dashboard.status_code, 200)
        self.assertEqual(app_dashboard.get_json()['total_plans'], 1)
        self.assertEqual(app_dashboard.get_json()['recent_plans'][0]['id'], plan_id)
        self.assertEqual(self.client.get('/api/apps/metadata').status_code, 200)
        self.assertEqual(APPSection.query.count(), 1)
        source_row = APPItem.query.one()
        self.assertEqual(source_row.source_ppmp_item_id, ppmp_item.id)
        self.assertEqual(source_row.project_title, ppmp_item.item_description)
        self.assertIsNone(source_row.early_procurement_activity)

        valid_source = {
            'section_id': source_row.section_id,
            'project_title': source_row.project_title,
            'implementing_unit': office.name,
            'general_description': 'Goods',
            'mode_of_procurement': 'Small Value Procurement',
            'early_procurement_activity': 'No',
            'bid_evaluation_criteria': 'LCRB',
            'procurement_start': '2026-01', 'procurement_end': '2026-12',
            'source_of_fund': 'GAA', 'estimated_budget': '1250.00',
            'procurement_strategy_tools': 'Documented procurement strategy',
            'remarks': 'Verified legitimate source row',
        }
        self.assertEqual(self.client.put(
            f'/api/apps/{plan_id}/items/{source_row.id}', json=valid_source,
        ).status_code, 200)
        section = self.client.post(
            f'/api/apps/{plan_id}/sections', json={'name': 'Capital Outlay'},
        )
        self.assertEqual(section.status_code, 201)
        second_item = dict(valid_source, section_id=section.get_json()['id'],
                           project_title='Authorized APP project', general_description='Infrastructure',
                           mode_of_procurement='Public Bidding', early_procurement_activity='Yes',
                           estimated_budget='500.00')
        second = self.client.post(f'/api/apps/{plan_id}/items', json=second_item)
        self.assertEqual(second.status_code, 201)
        invalid_timeline = dict(second_item, project_title='Invalid timeline',
                                procurement_start='2026-12', procurement_end='2026-01')
        self.assertEqual(self.client.post(
            f'/api/apps/{plan_id}/items', json=invalid_timeline,
        ).status_code, 400)
        self.assertEqual(self.client.post(
            f'/api/apps/{plan_id}/comments', json={'message': 'Permanent APP note'},
        ).status_code, 201)

        detail = self.client.get(f'/api/apps/{plan_id}').get_json()
        self.assertEqual(detail['item_count'], 2)
        self.assertEqual(detail['totals']['overall'], 1750.0)
        self.assertEqual(detail['totals']['early_procurement'], 500.0)
        self.assertEqual(detail['totals']['capital_outlay'], 500.0)
        self.assertEqual(APPComment.query.count(), 1)
        submitted = self.client.post(
            f'/api/apps/{plan_id}/transition', json={'action': 'submit'},
        )
        self.assertEqual(submitted.status_code, 200)
        self.assertEqual(submitted.get_json()['expected_reviewer'], 'BAC CHAIRMAN')
        self.assertEqual(self.client.put(
            f'/api/apps/{plan_id}/items/{source_row.id}', json=valid_source,
        ).status_code, 403)

        for reviewer, expected_status in (
            (bac, 'Under Review'), (eo, 'Under Review'), (budget, 'For Admin Approval'),
        ):
            with self.client.session_transaction() as session:
                session['user_id'], session['role'] = reviewer.id, reviewer.role
            reviewed = self.client.post(
                f'/api/apps/{plan_id}/transition', json={'action': 'approve'},
            )
            self.assertEqual(reviewed.status_code, 200)
            self.assertEqual(reviewed.get_json()['status'], expected_status)

        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        approved = self.client.post(
            f'/api/apps/{plan_id}/transition', json={'action': 'approve'},
        )
        self.assertEqual(approved.get_json()['status'], 'Approved')
        pdf = self.client.get(f'/api/apps/{plan_id}/export/pdf')
        excel = self.client.get(f'/api/apps/{plan_id}/export/excel')
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf.mimetype, 'application/pdf')
        self.assertEqual(excel.status_code, 200)
        self.assertEqual(excel.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        finalized = self.client.post(
            f'/api/apps/{plan_id}/transition', json={'action': 'finalize'},
        )
        self.assertEqual(finalized.get_json()['status'], 'Finalized')
        self.assertEqual(PPMPItem.query.count(), 1)
        self.assertEqual(ConsolidationSource.query.count(), 1)

    def test_pending_and_rejected_accounts_cannot_sign_in(self):
        office = College(code='REJECT-OFFICE', name='Reject Office')
        pending = User(
            username='pending_office', full_name='Pending Office', email='pending@isu.edu.ph',
            password_hash=generate_password_hash('password123'), role='REJECT-OFFICE',
            office='REJECT-OFFICE', college=office, account_status='PENDING', is_active=False,
        )
        db.session.add_all([office, pending])
        db.session.commit()
        login_payload = {'username': pending.username, 'password': 'password123'}
        self.assertIn(self.client.post('/api/auth/login', json=login_payload).status_code, {401, 403})

        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        rejected = self.client.post(
            f'/api/admin/users/{pending.id}/status', json={'status': 'REJECTED'},
        )
        self.assertEqual(rejected.status_code, 200)
        self.assertEqual(rejected.get_json()['user']['account_status'], 'REJECTED')
        self.assertIsNone(rejected.get_json()['user']['role'])
        self.client.post('/api/auth/logout')
        self.assertIn(self.client.post('/api/auth/login', json=login_payload).status_code, {401, 403})

    def test_legacy_wishlist_request_migration_is_idempotent_and_preserves_source(self):
        office = College(code='LEGACY', name='Legacy Office')
        requester = User(
            username='legacy_user', full_name='Legacy User', email='legacy@isu.edu.ph',
            password_hash='x', role='LEGACY', office='LEGACY', college=office,
            account_status='ACTIVE', is_active=True,
        )
        db.session.add_all([office, requester])
        db.session.flush()
        legacy_item = WishlistItem(
            office='LEGACY', created_by=requester.id, item_description='Legacy custom item',
            category='Other', unit='piece', quantity=12, estimated_cost=99,
            status='Pending Admin Approval',
        )
        db.session.add(legacy_item)
        db.session.flush()
        legacy_request = WishlistRequest(
            wishlist_item_id=legacy_item.id, requested_by=requester.id, quantity=4,
            remarks='Preserve this request', status='Pending Admin Approval',
        )
        db.session.add(legacy_request)
        db.session.commit()

        application.migrate_legacy_wishlist_orders()
        application.migrate_legacy_wishlist_orders()

        self.assertEqual(PPMPOrder.query.count(), 1)
        self.assertEqual(PPMPOrderItem.query.count(), 1)
        migrated = PPMPOrder.query.one()
        migrated_item = PPMPOrderItem.query.one()
        self.assertEqual(migrated.legacy_wishlist_request_id, legacy_request.id)
        self.assertEqual(migrated.reference_number, f'LEGACY-WR-{legacy_request.id:08d}')
        self.assertEqual(migrated.status, 'Pending Admin Approval')
        self.assertEqual(migrated_item.item_name, legacy_item.item_description)
        self.assertTrue(migrated_item.is_custom_item)
        self.assertIsNone(migrated_item.master_item_id)
        self.assertEqual(WishlistItem.query.count(), 1)

    def test_ppmp_order_reports_follow_role_permissions(self):
        budget = User(
            username='budget_report', full_name='Budget Report', email='budget-report@isu.edu.ph',
            password_hash='x', role='BUDGET', office='BUDGET',
            account_status='ACTIVE', is_active=True,
        )
        db.session.add(budget)
        db.session.commit()
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = budget.id, budget.role
        self.assertEqual(self.client.get('/api/reports/ppmp-orders').status_code, 200)
        self.assertEqual(self.client.get('/api/reports/ppmp-orders/export/pdf').status_code, 403)

    def test_admin_ppmp_and_app_dashboards_are_separate(self):
        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'] = admin.id
            session['role'] = admin.role
        ppmp = self.client.get('/api/dashboard/ppmp')
        app_dashboard = self.client.get('/api/dashboard/app')
        self.assertEqual(ppmp.status_code, 200)
        self.assertEqual(app_dashboard.status_code, 200)
        self.assertIn('eligible_for_consolidation', ppmp.get_json())
        self.assertIn('order_pending', ppmp.get_json())
        self.assertNotIn('total_amount', ppmp.get_json())
        self.assertNotIn('app_cse_count', ppmp.get_json())
        self.assertIn('app_cse_count', app_dashboard.get_json())
        self.assertNotIn('eligible_for_consolidation', app_dashboard.get_json())

    def test_removed_ppmp_dashboard_budget_components_preserve_shared_features(self):
        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'] = admin.id
            session['role'] = admin.role

        dashboard_html = self.client.get('/').get_data(as_text=True)
        for removed_text in (
            'Budget by Office', 'Budget by Fund Source', 'Budget by Procurement Type',
            'Budget by Quarter', 'Total Procurement Amount', 'id="stat-total-amount"',
            'id="chart-college"', 'id="chart-fund"', 'id="chart-type"',
            'id="chart-quarter"',
        ):
            self.assertNotIn(removed_text, dashboard_html)
        self.assertNotIn('function renderDashboard(', dashboard_html)

        # The revised Reports and Wishlist workspaces must remain intact.
        for preserved_text in (
            'report-status-donut', 'report-source-donut', 'report-office-bars',
            'report-item-bars', 'report-order-trend', 'report-records',
            'wishlist-category', 'wishlist-cost', 'Estimated Unit Cost',
            'ppmp-request-list', 'loadPPMPRequests',
        ):
            self.assertIn(preserved_text, dashboard_html)

        summary = self.client.get('/api/summary')
        self.assertEqual(summary.status_code, 200)
        self.assertNotIn('by_type', summary.get_json())
        for preserved_key in ('total_amount', 'by_college', 'by_fund', 'by_quarter'):
            self.assertIn(preserved_key, summary.get_json())

    def test_safe_dashboard_removal_does_not_break_preserved_workflows_or_exports(self):
        admin = User.query.filter_by(role='ADMIN').one()
        office = College(code='REMOVAL-SMOKE', name='Removal Smoke Office')
        db.session.add(office)
        db.session.flush()
        procurement = User(
            username='removal_procurement', full_name='Procurement',
            email='removal-procurement@isu.edu.ph', password_hash='x', role='PROCUREMENT',
            office='PROCUREMENT', account_status='ACTIVE', is_active=True,
        )
        db.session.add(procurement)
        db.session.flush()
        item = PPMPItem(
            college_id=office.id, fiscal_year=2026, item_description='Regression item',
            unit='piece', quantity=2, unit_cost=125, fund_source='GAA',
            mode_of_procurement='Public Bidding', procurement_type='Goods',
            quarter='Q1 (Jan-Mar)', status='Approved', created_by=admin.id,
            approved_by=admin.id,
        )
        db.session.add(item)
        db.session.commit()
        with self.client.session_transaction() as session:
            session['user_id'] = admin.id
            session['role'] = admin.role

        wishlist = self.client.post('/api/wishlist', json={
            'item_description': 'Wishlist regression item', 'category': 'Office Supplies',
            'unit': 'box', 'estimated_cost': 500,
        })
        self.assertEqual(wishlist.status_code, 201)
        self.assertEqual(wishlist.get_json()['category'], 'Office Supplies')
        self.assertEqual(wishlist.get_json()['estimated_cost'], 500)

        consolidated = self.client.get('/api/consolidated')
        self.assertEqual(consolidated.status_code, 200)
        office_summary = next(row for row in consolidated.get_json() if row['college_id'] == office.id)
        self.assertEqual(office_summary['by_fund']['GAA'], 250)
        batch = self.client.post('/api/consolidation-batches', json={
            'fiscal_year': 2026, 'ppmp_item_ids': [item.id], 'app_cycle': 'ANNUAL',
        })
        self.assertEqual(batch.status_code, 201)
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = procurement.id, procurement.role
        plan = self.client.post('/api/apps', json={
            'office': office.code, 'fiscal_year': 2026, 'plan_type': 'APP',
            'consolidation_batch_id': batch.get_json()['id'],
        })
        self.assertEqual(plan.status_code, 201)
        self.assertEqual(self.client.get('/api/dashboard/app').status_code, 200)

        pdf = self.client.get('/api/export/pdf')
        excel = self.client.get('/api/export/excel')
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf.mimetype, 'application/pdf')
        self.assertEqual(excel.status_code, 200)
        self.assertEqual(excel.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        anonymous = application.app.test_client()
        self.assertEqual(anonymous.get('/api/export/pdf').status_code, 401)
        self.assertEqual(anonymous.get('/api/consolidated').status_code, 401)

    def test_office_management_is_not_rendered_and_ppmp_wishlist_is_primary(self):
        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'] = admin.id
            session['role'] = admin.role
        admin_html = self.client.get('/').get_data(as_text=True)
        self.assertNotIn('<span class="icon">OF</span> Offices', admin_html)
        self.assertNotIn('Management Offices', admin_html)
        self.assertNotIn('Office Directory', admin_html)
        self.assertNotIn('panel-office-management', admin_html)
        self.assertNotIn('Managed Offices', admin_html)
        self.assertNotIn('Add Office', admin_html)
        self.assertIn('data-panel="ppmp-planning"', admin_html)
        for required_column in (
            '1. General Description / Objective', '2. Project Type', '3. Quantity / Size',
            '4. Procurement Method', '5. Pre-procurement Conference', '6. Start',
            '7. End', '8. Delivery / Implementation', '9. Source of Funds',
            '10. Estimated Budget', '11. Supporting Documents', '12. Remarks',
        ):
            self.assertIn(required_column, admin_html)
        self.assertIn('id="ppmp-item-method"></select>', admin_html)
        self.assertIn('id="ppmp-item-preproc"><option value="">Select Yes or No</option>', admin_html)
        self.assertIn('panel-dashboard', admin_html)
        self.assertIn('panel-ppmp-dashboard', admin_html)
        self.assertIn('panel-app-dashboard', admin_html)
        self.assertIn('id="ppmp-dashboard-nav"', admin_html)
        self.assertIn('id="app-dashboard-nav"', admin_html)
        self.assertIn('id="ppmp-master-nav"', admin_html)
        self.assertNotIn('id="consolidation-nav"', admin_html)
        self.assertIn('PPMP Wishlist', admin_html)
        self.assertIn('Account Approval', admin_html)
        self.assertNotIn('Accounts & Recovery', admin_html)
        self.assertIn('<section class="panel" id="panel-ppmp-wishlist">', admin_html)
        self.assertNotIn('panel-ppmp-wishlist-legacy', admin_html)
        self.assertIn('id="panel-ppmp-planning"', admin_html)
        self.assertIn('id="app-order-consolidation-card"', admin_html)
        self.assertIn('id="admin-wishlist-form-card"', admin_html)
        self.assertIn('Add Item to CSE Master List', admin_html)
        self.assertIn('<h1>CSE Master List</h1>', admin_html)
        self.assertIn('Official Wishlist catalog for PPMP. Only the Administrator can edit or delete these entries.', admin_html)
        self.assertIn('onclick="editWishlistItem(${item.id})"', admin_html)
        self.assertIn('onclick="deleteWishlistItem(${item.id},this)"', admin_html)
        self.assertNotIn('Add PPMP Planning Row', admin_html)
        self.assertNotIn('onclick="showPanel(\'ppmp-planning\')">PPMP Planning Documents</button>', admin_html)
        self.assertNotIn('Add Master Wishlist Item', admin_html)
        self.assertNotIn('Master Wishlist', admin_html)
        self.assertNotIn('Current PPMP Order', admin_html)
        self.assertNotIn('Add Unlisted Item', admin_html)
        self.assertNotIn('Available Quantity', admin_html)
        self.assertNotIn('id="wishlist-quantity"', admin_html)
        self.assertIn('<h1>REPORTS</h1>', admin_html)
        self.assertIn('Available Reports', admin_html)
        self.assertIn('Updated Reports from Administrator', admin_html)
        self.assertIn('Generate Report', admin_html)
        self.assertIn('Export PDF', admin_html)
        self.assertIn('Export Excel', admin_html)
        self.assertIn('id="ppmp-request-report-detail"', admin_html)
        self.assertIn('PROJECT PROCUREMENT MANAGEMENT PLAN (PPMP) NO. 1', admin_html)
        self.assertIn('General Description and Objective of the Project to be Procured', admin_html)
        self.assertIn('Expected Delivery Implementation Period', admin_html)
        self.assertIn('ATTACHED SUPPORTING DOCUMENTS', admin_html)
        self.assertIn('>Export <span aria-hidden="true">▼</span></button>', admin_html)
        self.assertIn("exportPPMPRequestRecord('pdf')\">PDF</button>", admin_html)
        self.assertIn("exportPPMPRequestRecord('excel')\">Excel</button>", admin_html)
        self.assertIn('>Print Preview</button>', admin_html)
        self.assertNotIn('Export to PDF or Excel', admin_html)
        self.assertNotIn('<th>Action</th></tr></thead><tbody id="ppmp-request-report-items"', admin_html)
        self.assertIn('id="wishlist-comment-modal"', admin_html)
        self.assertNotIn("prompt('Enter an Administrator comment", admin_html)
        for page_html in (admin_html,):
            self.assertIn('<span class="icon">PP</span> PPMP Dashboard', page_html)
            self.assertIn('<span class="icon">WL</span> PPMP Wishlist', page_html)
            self.assertIn('<span class="icon">WL</span> CSE Master List', page_html)
            self.assertIn('<span class="icon">RQ</span> PPMP Requests', page_html)
            self.assertIn('Open PPMP Items', page_html)
            self.assertNotIn('Open PPMP Items &amp; Orders', page_html)
            self.assertNotIn('Open PPMP Item &amp; Orders', page_html)
            self.assertIn('<div class="stat-label">Pending Requests</div>', page_html)
            self.assertNotIn('<div class="stat-label">Pending Orders</div>', page_html)
            self.assertIn('<div class="stat-label">Approved Requests</div>', page_html)
            self.assertNotIn('<div class="stat-label">Approved Items</div>', page_html)

        user = User(username='budget', full_name='Budget', email='budget@isu.edu.ph', password_hash='x',
                    role='BUDGET', office='BUDGET', account_status='ACTIVE', is_active=True)
        db.session.add(user)
        db.session.commit()
        with self.client.session_transaction() as session:
            session['user_id'] = user.id
            session['role'] = user.role
        user_html = self.client.get('/').get_data(as_text=True)
        self.assertIn('id="order-year"', user_html)
        self.assertIn('aria-label="Fiscal Year" placeholder="Input Fiscal Year"', user_html)
        self.assertIn('id="order-quarter"', user_html)
        self.assertNotIn('id="order-fund"', user_html)
        self.assertNotIn('id="order-mode"', user_html)
        self.assertNotIn('id="order-type"', user_html)
        self.assertNotIn('id="order-remarks"', user_html)
        self.assertNotIn("populate('order-fund'", user_html)
        self.assertNotIn("populate('order-mode'", user_html)
        self.assertNotIn("populate('order-type'", user_html)
        self.assertIn('data-report-field="project_type" data-report-required', user_html)
        self.assertIn('data-report-field="source_of_funds" data-report-required', user_html)
        self.assertIn('data-report-field="remarks" maxlength="5000"', user_html)
        self.assertNotRegex(user_html, r'id="panel-office-management"')
        self.assertNotRegex(user_html, r'id="office-management-nav"')
        self.assertRegex(user_html, r'id="ppmp-dashboard-nav"')
        self.assertRegex(user_html, r'id="app-dashboard-nav"')
        self.assertRegex(user_html, r'id="consolidation-nav"')
        self.assertNotIn('id="admin-wishlist-form-card"', user_html)
        self.assertNotIn('id="wishlist-description"', user_html)
        self.assertNotIn('id="wishlist-comment-modal"', user_html)
        self.assertIn('Current PPMP Items', user_html)
        self.assertIn('Add other items not included in the CSE Master list', user_html)
        self.assertIn('onclick="addCustomOrderItem(this)">Add Item</button>', user_html)
        self.assertNotIn('onclick="addCustomOrderItem(this)">Add other items not included in the CSE Master list</button>', user_html)
        self.assertIn('Open PPMP Items', user_html)
        self.assertIn('<div class="stat-label">Pending Requests</div>', user_html)
        self.assertIn('<div class="stat-label">Approved Requests</div>', user_html)
        self.assertNotIn('<span class="icon">PP</span> PPMP Dashboard', user_html)
        self.assertNotIn('<span class="icon">WL</span> PPMP Wishlist', user_html)
        self.assertNotIn('<span class="icon">WL</span> CSE Master List', user_html)
        self.assertNotIn('<span class="icon">RQ</span> PPMP Requests', user_html)
        self.assertRegex(user_html, r'id="custom-item-category"[^>]*>\s*<option value="">Select Category</option><option value="Goods">Goods</option><option value="Infrastructure">Infrastructure</option><option value="Consulting Services">Consulting Services</option>')

        definitions = set(re.findall(r'(?:async\s+)?function\s+([A-Za-z_$][\w$]*)\s*\(', admin_html))
        expected_handlers = {
            'loadPPMPDashboard',
            'loadAPPDashboard', 'loadAPPConsolidationQueue', 'loadAdminOverview', 'ppmpItemActions', 'transitionPPMP',
            'appPlanActions', 'transitionAPP', 'refreshAccountAdministration',
            'loadWishlist', 'createWishlistItem', 'editWishlistItem', 'deleteWishlistItem',
            'queueWishlistSearch', 'queueDashboardWishlistSearch',
            'resetWishlistItemForm', 'loadPPMPRequests', 'loadPPMPPlanning',
            'transitionPPMPOrder', 'addMasterOrderItem', 'addCustomOrderItem',
            'loadDashboardWishlist', 'openCustomPPMPItem',
            'addWishlistOrderComment', 'closeWishlistCommentModal', 'saveWishlistOrderComment',
            'submitPPMPOrder', 'submitWishlistRequest', 'loadOrderReports',
            'loadAvailableReports', 'generateAuthorizedReport', 'reviewAvailableReport',
            'exportAuthorizedReport', 'saveAvailableReport', 'editAvailableReport',
            'resetAvailableReportForm', 'openPPMPRequestReport',
            'savePPMPRequestReportItem', 'exportPPMPRequestRecord',
            'toggleProfile', 'loadActivityLog', 'toggleActivityTracker',
            'loadAdminActivityTracker',
        }
        self.assertEqual(expected_handlers - definitions, set(), 'The revised UI handlers must all be implemented.')

    def test_revision_migration_is_idempotent(self):
        application.migrate_revision_schema()
        application.migrate_revision_schema()
        inspector = application.inspect(db.engine)
        order_columns = {column['name'] for column in inspector.get_columns('ppmp_orders')}
        item_columns = {column['name'] for column in inspector.get_columns('ppmp_order_items')}
        item_column_map = {column['name']: column for column in inspector.get_columns('ppmp_order_items')}
        self.assertIn('ppmp_reports', inspector.get_table_names())
        self.assertIn('user_activity_logs', inspector.get_table_names())
        activity_columns = {column['name'] for column in inspector.get_columns('user_activity_logs')}
        self.assertTrue({'user_id', 'time_in', 'time_out'} <= activity_columns)
        self.assertTrue({'fiscal_year', 'fund_source', 'mode_of_procurement', 'procurement_type', 'quarter'} <= order_columns)
        self.assertTrue({'master_item_id', 'category', 'estimated_unit_cost', 'is_custom_item'} <= item_columns)
        self.assertEqual(item_column_map['procurement_method']['type'].length, 255)

    def test_global_loading_overlay_and_duplicate_action_guards_are_present(self):
        login_response = application.app.test_client().get('/')
        login_html = login_response.get_data(as_text=True)
        self.assertEqual(login_response.headers['Cache-Control'], 'no-store, no-cache, must-revalidate, max-age=0')
        self.assertIn('content="infinity-auth-v2-20260910"', login_html)
        self.assertIn('class="infinity-loop"', login_html)
        self.assertIn('const AUTH_LOADING_DURATION_MS = 10_000;', login_html)
        self.assertNotIn('const AUTH_LOADING_DURATION_MS = 15_000;', login_html)
        self.assertNotIn('const AUTH_LOADING_DURATION_MS = 30_000;', login_html)
        self.assertIn('await authLoading.wait();location.assign', login_html)
        self.assertIn('authLoading.stop();', login_html)
        self.assertIn('capTimer=setTimeout(hideAuthLoading,AUTH_LOADING_DURATION_MS)', login_html)
        self.assertIn('function hideAuthLoading()', login_html)
        self.assertNotIn('Secure transition', login_html)
        self.assertNotIn('id="loading-duration"', login_html)
        self.assertNotIn('import { InfinityLoop }', login_html)

        admin = User.query.filter_by(role='ADMIN').one()
        with self.client.session_transaction() as session:
            session['user_id'], session['role'] = admin.id, admin.role
        html = self.client.get('/').get_data(as_text=True)
        self.assertIn('data-websocket-url=""', html)
        self.assertIn('data-is-admin="true"', html)
        self.assertIn('const appConfig = document.body;', html)
        self.assertIn("const WS = appConfig.dataset.websocketUrl || '';", html)
        self.assertIn("const isAdminSession = appConfig.dataset.isAdmin === 'true';", html)
        self.assertNotIn('{{ websocket_url', html)
        self.assertIn('content="infinity-auth-v2-20260910"', html)
        self.assertIn('.app-loading{position:fixed;inset:0;z-index:5000', html)
        self.assertIn('id="app-loading"', html)
        self.assertIn('id="app-loading-message"', html)
        self.assertIn('class="app-loading" id="app-loading"', html)
        self.assertNotIn('class="app-loading visible" id="app-loading"', html)
        self.assertIn('aria-busy="false"', html)
        self.assertNotIn('id="app-loading-duration"', html)
        self.assertNotIn('Secure transition', html)
        self.assertIn('class="infinity-loop"', html)
        self.assertNotIn('const AUTH_LOADING_DURATION_MS', html)
        self.assertNotIn('const AUTH_LOADING_DURATION_MS = 30_000;', html)
        self.assertIn("await apiFetch('/api/auth/logout', { method: 'POST', showLoading: false });", html)
        self.assertNotIn('beginTimedAuthLoading', html)
        self.assertNotIn('await authLoading.wait();', html)
        self.assertNotIn("message.textContent = 'Signing you out", html)
        self.assertIn('setTimeout(() => controller.abort(), 30000)', html)
        self.assertIn("showLoading = requestMethod !== 'GET'", html)
        self.assertIn('async function loadItems(showLoading = false)', html)
        self.assertIn('async function loadWishlist(showLoading = false)', html)
        self.assertIn('async function loadDashboardWishlist(showLoading = false)', html)
        self.assertIn('async function generateAuthorizedReport(showLoading = false', html)
        self.assertIn('async function openAPP(id, showLoading=false)', html)
        self.assertIn('const SEARCH_DEBOUNCE_MS = 250;', html)
        self.assertIn('const LOADING_FAILSAFE_MS = 31_000;', html)
        self.assertIn('setTimeout(resetApplicationLoading, LOADING_FAILSAFE_MS)', html)
        self.assertIn('oninput="queueWishlistSearch()"', html)
        self.assertIn('oninput="queueDashboardWishlistSearch()"', html)
        self.assertIn('loadWishlist(false)', html)
        self.assertIn('loadDashboardWishlist(false)', html)
        self.assertIn('function resetApplicationLoading()', html)
        self.assertIn("method: 'POST', body: JSON.stringify(payload), loadingMessage: 'Submitting PPMP request...'", html)
        self.assertIn('finally {', html)
        self.assertIn('button.disabled = true', html)
        self.assertIn("button.textContent = 'Processing...'", html)


if __name__ == '__main__':
    unittest.main()
