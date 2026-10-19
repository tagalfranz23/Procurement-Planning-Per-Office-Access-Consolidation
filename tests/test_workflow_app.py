"""Integration coverage for official-source consolidation and the agreed APP flow.

All data lives in SQLite memory. Never import this module with the production
database already bound to the application.
"""
import io
import os
import unittest
from datetime import date
from decimal import Decimal

os.environ['DATABASE_URL'] = 'sqlite:///:memory:'
os.environ['FLASK_SECRET_KEY'] = 'workflow-app-tests-only'

import app as application
from models import (
    APPApproval, APPItem, APPItemContribution, AnnualProcurementPlan, College,
    ConsolidationBatch, ConsolidationSource, OfficialPPMP, PPMPItem, PPMPOrder,
    PPMPOrderItem, User, WishlistItem, WorkflowAuditLog, db,
)


class ConsolidatedAPPWorkflowTest(unittest.TestCase):
    def setUp(self):
        application.app.config.update(TESTING=True)
        self.context = application.app.app_context()
        self.context.push()
        self.assertEqual(db.engine.url.database, ':memory:', 'These tests may only use an in-memory database.')
        db.drop_all()
        db.create_all()
        self.users = {}
        for role in ('ADMIN', 'PROCUREMENT', 'BAC CHAIRMAN', 'BUDGET', 'EO', 'CBAO', 'CASHIER'):
            college = None
            if role in {'CBAO', 'CASHIER'}:
                college = College(code=role, name=f'{role} Office', is_active=True)
                db.session.add(college)
                db.session.flush()
            user = User(
                username=role.replace(' ', '_'), full_name=f'{role} User',
                email=f'{role.replace(" ", "_")}@example.test', password_hash='unused',
                role=role, office=role, college_id=college.id if college else None,
                account_status='ACTIVE', is_active=True,
            )
            db.session.add(user)
            self.users[role] = user
        db.session.commit()
        self.client = application.app.test_client()
        self.login('ADMIN')

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def login(self, role):
        user = self.users[role]
        with self.client.session_transaction() as session:
            session['user_id'] = user.id
            session['role'] = role
            session['auth_version'] = user.auth_version or 0

    def source(self, office='CBAO', *, title='Printer paper', quantity='2.00', budget='120.25',
               year=2027, fund='GAA/101', mode='Small Value Procurement', unit='reams',
               specifications='A4, 80 gsm', start=None, end=None, delivery=None, needed=False,
               official=True, is_cse=False):
        user = self.users[office]
        order = PPMPOrder.query.filter_by(office_id=user.college_id, fiscal_year=year, workflow_official=official).first()
        if order is None:
            order = PPMPOrder(
                reference_number=f'PPMP-{office}-{year}-{PPMPOrder.query.count() + 1}',
                office_id=user.college_id, requested_by=user.id, fiscal_year=year,
                workflow_official=official, status='Approved',
            )
            db.session.add(order)
            db.session.flush()
            if official:
                db.session.add(OfficialPPMP(office_id=user.college_id, fiscal_year=year, source_order_id=order.id))
        ppmp = PPMPItem(
            college_id=user.college_id, fiscal_year=year, item_description=title, unit=unit,
            quantity=Decimal(quantity), unit_cost=Decimal('1.00'), authorized_budget=Decimal(budget),
            fund_source=fund, mode_of_procurement=mode, procurement_type='Goods/Services',
            status='Approved', created_by=user.id, approved_by=self.users['ADMIN'].id,
        )
        db.session.add(ppmp)
        db.session.flush()
        master = None
        if is_cse:
            master = WishlistItem(item_description=title, category='Goods', unit=unit,
                                  quantity=Decimal('1.00'), estimated_cost=Decimal('1.00'), status='Available')
            db.session.add(master)
            db.session.flush()
        db.session.add(PPMPOrderItem(
            order=order, ppmp_item_id=ppmp.id, master_item_id=master.id if master else None,
            item_name=title, description=specifications,
            project_type='Goods/Services', quantity=Decimal(quantity), unit=unit,
            procurement_method=mode, pre_procurement_conference=False,
            procurement_start=None if needed else start or date(year, 1, 1),
            procurement_end=None if needed else end or date(year, 2, 1),
            delivery_period=None if needed else delivery or date(year, 3, 1),
            schedule_as_needed=needed, source_of_funds=fund,
            estimated_budget=Decimal(budget), is_custom_item=not is_cse, created_by=user.id,
        ))
        db.session.commit()
        return ppmp

    def batch(self, sources, year=2027):
        self.login('ADMIN')
        response = self.client.post('/api/consolidation-batches', json={
            'fiscal_year': year, 'ppmp_item_ids': [source.id for source in sources],
        })
        self.assertEqual(response.status_code, 201, response.get_json())
        return response.get_json()

    def plan(self, sources):
        batch = self.batch(sources, sources[0].fiscal_year)
        response = self.client.post('/api/apps', json={'batch_id': batch['id']})
        self.assertEqual(response.status_code, 201, response.get_json())
        return response.get_json()

    def complete(self, plan_id):
        self.login('ADMIN')
        plan = db.session.get(AnnualProcurementPlan, plan_id)
        for item in list(plan.items):
            response = self.client.put(f'/api/apps/{plan_id}/items/{item.id}', json={
                'early_procurement_activity': False,
                'bid_evaluation_criteria': 'Lowest calculated responsive bid',
            })
            self.assertEqual(response.status_code, 200, response.get_json())

    def transition(self, plan_id, role, action, note=None, expected=200):
        self.login(role)
        response = self.client.post(f'/api/apps/{plan_id}/transition', json={'action': action, 'note': note})
        self.assertEqual(response.status_code, expected, response.get_json())
        return response.get_json()

    def test_admin_only_and_official_same_year_sources(self):
        source = self.source()
        legacy = self.source('CASHIER', official=False)
        self.login('PROCUREMENT')
        self.assertEqual(self.client.get('/api/consolidation-eligible').status_code, 403)
        self.assertEqual(self.client.post('/api/consolidation-batches', json={
            'fiscal_year': 2027, 'ppmp_item_ids': [source.id],
        }).status_code, 403)
        self.login('ADMIN')
        eligible = self.client.get('/api/consolidation-eligible').get_json()
        self.assertEqual([row['source_ppmp_item_id'] for row in eligible['items']], [source.id])
        for ids, year in (([source.id], 2028), ([legacy.id], 2027), (['1'], 2027)):
            response = self.client.post('/api/consolidation-batches', json={'fiscal_year': year, 'ppmp_item_ids': ids})
            self.assertIn(response.status_code, (400, 409))
        self.assertEqual(ConsolidationBatch.query.count(), 0)
        batch = self.batch([source])
        self.assertEqual(self.client.post('/api/consolidation-batches', json={
            'fiscal_year': 2027, 'ppmp_item_ids': [source.id],
        }).status_code, 409)
        self.assertEqual(ConsolidationSource.query.count(), 1)
        self.assertEqual(batch['offices'], ['CBAO'])

    def test_strict_merging_keeps_exact_original_office_contributions(self):
        first = self.source(quantity='1.25', budget='100.01')
        second = self.source('CASHIER', quantity='2.75', budget='200.02')
        different = [
            self.source(fund='SB 164'), self.source(unit='boxes'),
            self.source(specifications='Legal, 70 gsm'),
            self.source(start=date(2027, 2, 1)), self.source(end=date(2027, 4, 1)),
            self.source(delivery=date(2027, 5, 1)), self.source(needed=True),
            self.source(mode='Competitive Bidding'),
        ]
        result = self.plan([first, second] + different)
        plan = db.session.get(AnnualProcurementPlan, result['id'])
        self.assertEqual(len(plan.items), 1 + len(different))
        merged = next(row for row in plan.items if len(row.contributions) == 2)
        self.assertEqual(merged.quantity, Decimal('4.00'))
        self.assertEqual(merged.estimated_budget, Decimal('300.03'))
        self.assertEqual({row.office.code for row in merged.contributions}, {'CBAO', 'CASHIER'})
        self.assertEqual(merged.contributions[0].source_snapshot['specifications'], 'A4, 80 gsm')
        self.assertEqual(first.authorized_budget, Decimal('100.01'))
        self.assertEqual(first.status, 'Approved')

    def test_multiple_unlinked_batches_attach_in_any_order_without_duplicate_sources(self):
        first, second = self.source(), self.source('CASHIER')
        first_batch = self.batch([first])
        second_batch = self.batch([second])
        self.assertNotEqual(first_batch['app_cycle'], second_batch['app_cycle'])
        initial = self.client.post('/api/apps', json={'batch_id': second_batch['id']})
        self.assertEqual(initial.status_code, 201, initial.get_json())
        appended = self.client.post('/api/apps', json={'batch_id': first_batch['id']})
        self.assertEqual(appended.status_code, 200, appended.get_json())
        self.assertEqual(initial.get_json()['id'], appended.get_json()['id'])
        self.assertEqual(AnnualProcurementPlan.query.count(), 1)
        self.assertEqual(APPItemContribution.query.count(), 2)
        self.assertEqual(len(appended.get_json()['sections'][0]['items']), 1)
        duplicate = self.client.post('/api/apps', json={'batch_id': second_batch['id']})
        self.assertEqual(duplicate.status_code, 409)
        self.assertEqual(APPItemContribution.query.count(), 2)

    def test_cse_total_uses_only_cse_office_allocation_in_mixed_merged_row(self):
        result = self.plan([self.source(is_cse=True, budget='80.05'), self.source('CASHIER', budget='40.00')])
        self.assertEqual(result['item_count'], 1)
        self.assertEqual(result['totals']['overall'], 120.05)
        self.assertEqual(result['totals']['cse'], 80.05)
        item = result['sections'][0]['items'][0]
        contributions = [
            {'id': row['id'], 'quantity': '1.00', 'budget': '50.00' if row['source_snapshot']['is_cse'] else '20.00'}
            for row in item['contributions']
        ]
        response = self.client.put(f'/api/apps/{result["id"]}/items/{item["id"]}/contributions',
                                   json={'contributions': contributions})
        self.assertEqual(response.status_code, 200, response.get_json())
        totals = db.session.get(AnnualProcurementPlan, result['id']).totals()
        self.assertEqual(totals['overall'], 70)
        self.assertEqual(totals['cse'], 50)

    def test_as_needed_timeline_stays_visible_in_app_exports(self):
        from openpyxl import load_workbook
        result = self.plan([self.source(needed=True)])
        self.complete(result['id'])
        self.assertIsNone(application.validate_app_for_submission(db.session.get(AnnualProcurementPlan, result['id'])))
        row = self.client.get(f'/api/apps/{result["id"]}').get_json()['sections'][0]['items'][0]
        self.assertTrue(row['schedule_as_needed'])
        response = self.client.get(f'/api/apps/{result["id"]}/export/excel')
        workbook = load_workbook(io.BytesIO(response.data))
        row = next(row for row in workbook.active.values if row[0] and str(row[0]).startswith('Printer paper'))
        self.assertEqual(row[6:8], ('As Needed', 'As Needed'))

    def test_review_order_finalizes_on_admin_approval_and_locks(self):
        result = self.plan([self.source()])
        plan_id = result['id']
        self.complete(plan_id)
        self.transition(plan_id, 'PROCUREMENT', 'submit', expected=403)
        submitted = self.transition(plan_id, 'ADMIN', 'submit')
        self.assertEqual(submitted['expected_reviewer'], 'BAC CHAIRMAN')
        self.transition(plan_id, 'EO', 'approve', expected=403)
        for role, next_role in (('BAC CHAIRMAN', 'BUDGET'), ('BUDGET', 'EO'), ('EO', 'ADMIN')):
            result = self.transition(plan_id, role, 'approve')
            self.assertEqual(result['expected_reviewer'], next_role)
        result = self.transition(plan_id, 'ADMIN', 'approve')
        self.assertEqual(result['status'], 'Finalized')
        self.assertIsNone(result['expected_reviewer'])
        item_id = result['sections'][0]['items'][0]['id']
        self.assertEqual(self.client.put(f'/api/apps/{plan_id}/items/{item_id}', json={'remarks': 'change'}).status_code, 403)
        approvals = APPApproval.query.filter_by(app_id=plan_id, action='APPROVE').all()
        self.assertEqual([row.reviewer_role for row in approvals], ['BAC CHAIRMAN', 'BUDGET', 'EO', 'ADMIN'])
        self.assertTrue(all(row.snapshot['rows'][0]['contributions'] for row in approvals))

    def test_returned_budget_review_resumes_and_revised_allocations_reconcile(self):
        sources = [self.source(), self.source('CASHIER')]
        result = self.plan(sources)
        plan_id = result['id']
        self.complete(plan_id)
        self.transition(plan_id, 'ADMIN', 'submit')
        self.transition(plan_id, 'BAC CHAIRMAN', 'approve')
        returned = self.transition(plan_id, 'BUDGET', 'return', 'Please revise office allocation')
        self.assertEqual(returned['return_review_role'], 'BUDGET')
        item = returned['sections'][0]['items'][0]
        self.login('ADMIN')
        contributions = [{'id': row['id'], 'quantity': '1.00', 'budget': '60.00'} for row in item['contributions']]
        invalid = self.client.put(f'/api/apps/{plan_id}/items/{item["id"]}/contributions', json={
            'contributions': contributions, 'quantity': '2.00', 'budget': '121.00',
        })
        self.assertEqual(invalid.status_code, 400)
        self.assertEqual(db.session.get(APPItem, item['id']).estimated_budget, Decimal('240.50'))
        revised = self.client.put(f'/api/apps/{plan_id}/items/{item["id"]}/contributions', json={
            'contributions': contributions, 'quantity': '2.00', 'budget': '120.00', 'note': 'Agreed office allocations',
        })
        self.assertEqual(revised.status_code, 200, revised.get_json())
        self.assertTrue(all(row.original_budget == Decimal('120.25') for row in APPItemContribution.query.all()))
        self.assertTrue(all(source.authorized_budget == Decimal('120.25') for source in sources))
        submitted = self.transition(plan_id, 'ADMIN', 'submit')
        self.assertEqual(submitted['expected_reviewer'], 'BUDGET')
        self.assertEqual(APPApproval.query.filter_by(app_id=plan_id, reviewer_role='BAC CHAIRMAN', action='APPROVE').count(), 1)

    def test_late_batch_preserves_stage_and_requires_new_row_completion(self):
        first = self.source()
        result = self.plan([first])
        plan_id = result['id']
        first_item_id = result['sections'][0]['items'][0]['id']
        self.complete(plan_id)
        self.transition(plan_id, 'ADMIN', 'submit')
        self.transition(plan_id, 'BAC CHAIRMAN', 'approve')
        late = self.source('CASHIER', title='Scanner')
        batch = self.batch([late])
        response = self.client.post('/api/apps', json={'batch_id': batch['id']})
        self.assertEqual(response.status_code, 200, response.get_json())
        result = response.get_json()
        self.assertEqual(result['id'], plan_id)
        self.assertEqual(result['expected_reviewer'], 'BUDGET')
        pending = [row for section in result['sections'] for row in section['items'] if row['is_pending_completion']]
        self.assertEqual(len(pending), 1)
        self.assertTrue(pending[0]['can_edit'])
        self.assertEqual(self.client.put(f'/api/apps/{plan_id}/items/{first_item_id}', json={'remarks': 'blocked'}).status_code, 403)
        self.transition(plan_id, 'BUDGET', 'approve', expected=409)
        self.login('ADMIN')
        response = self.client.put(f'/api/apps/{plan_id}/items/{pending[0]["id"]}', json={
            'early_procurement_activity': True, 'bid_evaluation_criteria': 'Responsive bid',
        })
        self.assertEqual(response.status_code, 200, response.get_json())
        approved = self.transition(plan_id, 'BUDGET', 'approve')
        self.assertEqual(approved['expected_reviewer'], 'EO')
        self.assertEqual(APPApproval.query.filter_by(app_id=plan_id, reviewer_role='BAC CHAIRMAN', action='APPROVE').count(), 1)
        history = WorkflowAuditLog.query.filter_by(entity_type='app', entity_id=plan_id, action='APP_BATCH_ADDED').one()
        self.assertEqual(history.details['snapshot']['item_count'], 2)
        self.assertEqual(len(history.details['snapshot']['batch_ids']), 2)

    def test_supplemental_app_preserves_final_annual(self):
        result = self.plan([self.source()])
        annual_id = result['id']
        self.complete(annual_id)
        self.transition(annual_id, 'ADMIN', 'submit')
        for role in ('BAC CHAIRMAN', 'BUDGET', 'EO', 'ADMIN'):
            self.transition(annual_id, role, 'approve')
        annual_revision = db.session.get(AnnualProcurementPlan, annual_id).revision
        supplemental = self.plan([self.source('CASHIER')])
        self.assertEqual(supplemental['supplemental_parent_id'], annual_id)
        self.assertEqual(supplemental['supplemental_number'], 1)
        annual = db.session.get(AnnualProcurementPlan, annual_id)
        self.assertEqual(annual.status, 'Finalized')
        self.assertEqual(annual.revision, annual_revision)
        self.assertEqual(len(annual.items), 1)

    def test_returned_late_row_completion_does_not_leave_submitted_edit_permission(self):
        result = self.plan([self.source()])
        plan_id = result['id']
        self.complete(plan_id)
        self.transition(plan_id, 'ADMIN', 'submit')
        self.transition(plan_id, 'BAC CHAIRMAN', 'approve')
        batch = self.batch([self.source('CASHIER', title='Late scanner')])
        appended = self.client.post('/api/apps', json={'batch_id': batch['id']}).get_json()
        added_item_id = appended['source_details']['pending_item_ids'][0]
        self.transition(plan_id, 'BUDGET', 'return', 'Complete the late scanner row')
        self.complete(plan_id)
        submitted = self.transition(plan_id, 'ADMIN', 'submit')
        self.assertEqual(submitted['expected_reviewer'], 'BUDGET')
        self.assertEqual(submitted['source_details'].get('pending_item_ids', []), [])
        denied = self.client.put(f'/api/apps/{plan_id}/items/{added_item_id}', json={'remarks': 'submitted edit'})
        self.assertEqual(denied.status_code, 403)

    def test_office_feedback_is_scoped_without_app_access(self):
        first, second = self.source(), self.source('CASHIER', title='Scanner')
        result = self.plan([first, second])
        order = PPMPOrderItem.query.filter_by(ppmp_item_id=first.id).one().order
        other_order = PPMPOrderItem.query.filter_by(ppmp_item_id=second.id).one().order
        self.login('CBAO')
        self.assertEqual(self.client.get(f'/api/apps/{result["id"]}').status_code, 403)
        feedback = self.client.get(f'/api/ppmp-orders/{order.id}/app-feedback')
        self.assertEqual(feedback.status_code, 200, feedback.get_json())
        allocations = feedback.get_json()['apps'][0]['contributions']
        self.assertEqual([row['source_ppmp_item_id'] for row in allocations], [first.id])
        self.assertEqual(self.client.get(f'/api/ppmp-orders/{other_order.id}/app-feedback').status_code, 403)
        response = self.client.post(f'/api/ppmp-orders/{order.id}/app-feedback', json={
            'app_id': result['id'], 'message': 'Our office confirms its quantity.',
        })
        self.assertEqual(response.status_code, 201, response.get_json())
        self.assertEqual(self.client.post(f'/api/ppmp-orders/{order.id}/app-feedback', json={
            'app_id': result['id'] + 100, 'message': 'Forbidden target',
        }).status_code, 400)

    def test_allocation_precision_and_ids_reject_before_write(self):
        result = self.plan([self.source()])
        item = result['sections'][0]['items'][0]
        allocation = item['contributions'][0]
        for quantity, budget, identity in (
            ('NaN', '10.00', allocation['id']), ('1.00', 'Infinity', allocation['id']),
            ('0.00', '10.00', allocation['id']), ('1.00', '10.001', allocation['id']),
            ('1.00', '10.00', allocation['id'] + 99),
        ):
            response = self.client.put(f'/api/apps/{result["id"]}/items/{item["id"]}/contributions', json={
                'contributions': [{'id': identity, 'quantity': quantity, 'budget': budget}],
            })
            self.assertEqual(response.status_code, 400, response.get_json())
        saved = db.session.get(APPItemContribution, allocation['id'])
        self.assertEqual(saved.quantity, Decimal('2.00'))
        self.assertEqual(saved.estimated_budget, Decimal('120.25'))

    def test_exports_preserve_twelve_columns_and_office_quantity_details(self):
        from openpyxl import load_workbook
        from PyPDF2 import PdfReader
        result = self.plan([self.source(), self.source('CASHIER')])
        response = self.client.get(f'/api/apps/{result["id"]}/export/excel')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data))
        sheet = workbook.active
        self.assertEqual(sheet.max_column, 12)
        values = [str(cell.value or '') for row in sheet for cell in row]
        self.assertTrue(any('Quantity: 4.00' in value for value in values))
        self.assertTrue(any('CBAO:' in value and 'CASHIER:' in value for value in values))
        pdf = self.client.get(f'/api/apps/{result["id"]}/export/pdf')
        self.assertEqual(pdf.status_code, 200)
        reader = PdfReader(io.BytesIO(pdf.data))
        self.assertGreaterEqual(len(reader.pages), 1)
        text = '\n'.join(page.extract_text() for page in reader.pages)
        self.assertIn('CBAO:', text)
        self.assertIn('CASHIER:', text)


if __name__ == '__main__':
    unittest.main()
