"""Official PPMP/copy workflow regressions using only an in-memory database."""
import io
import os
import unittest
from copy import deepcopy
from decimal import Decimal

# Set these before importing the Flask application; .env must never select a
# user's database when this test module is run independently.
os.environ['DATABASE_URL'] = 'sqlite:///:memory:'
os.environ['FLASK_SECRET_KEY'] = 'isolated-ppmp-workflow-test-secret'

from openpyxl import load_workbook
from flask import Flask
from sqlalchemy import inspect, text

import app as application
from models import (
    College, IndicativePPMPCopy, OfficialPPMP, PPMPItem, PPMPOrder,
    PPMPOrderItem, User, WishlistItem, WorkflowAuditLog, db,
)
from workflow_migrations import (
    WORKFLOW_REQUIRED_COLUMNS, WORKFLOW_REQUIRED_TABLES, apply_workflow_migrations,
)


class OfficialPPMPWorkflowTest(unittest.TestCase):
    def setUp(self):
        application.app.config.update(TESTING=True)
        self.context = application.app.app_context()
        self.context.push()
        self.assertEqual(db.engine.url.get_backend_name(), 'sqlite')
        self.assertEqual(db.engine.url.database, ':memory:')
        db.drop_all()
        db.create_all()
        self.office = College(code='CASHIER', name='Cashier Office')
        self.other_office = College(code='CBAO', name='Business Affairs Office')
        db.session.add_all([self.office, self.other_office])
        db.session.flush()
        self.owner = self.user('cashier', 'CASHIER', self.office)
        self.colleague = self.user('colleague', 'CAS', self.office)
        self.other = self.user('other', 'CBAO', self.other_office)
        self.procurement = self.user('procurement', 'PROCUREMENT')
        self.admin = self.user('admin', 'ADMIN')
        self.eo = self.user('eo', 'EO')
        self.master = WishlistItem(
            office='SUPPLY', item_description='Bond Paper A4', category='Goods',
            unit='ream', quantity=0, estimated_cost=Decimal('12.50'),
            status='Available', created_by=self.admin.id,
        )
        db.session.add(self.master)
        db.session.commit()
        self.client = application.app.test_client()
        self.login(self.owner)

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def user(self, username, role, office=None):
        user = User(
            username=username, full_name=username.title(), email=f'{username}@test.local',
            password_hash='unused-session-fixture', role=role,
            office=office.code if office else role,
            college_id=office.id if office else None,
            account_status='ACTIVE', is_active=True,
        )
        db.session.add(user)
        db.session.flush()
        return user

    def login(self, user):
        with self.client.session_transaction() as session:
            session.clear()
            session['user_id'] = user.id
            session['role'] = user.role
            session['auth_version'] = int(user.auth_version or 0)

    def draft_payload(self, fiscal_year=2028):
        return {
            'fiscal_year': fiscal_year, 'quarter': application.QUARTERS[0],
            'items': [
                {'master_item_id': self.master.id, 'quantity': '2'},
                {'item_name': 'Office measuring tool', 'category': 'Goods',
                 'unit': 'piece', 'quantity': '1', 'estimated_unit_cost': '100.00'},
            ],
        }

    def create_draft(self, fiscal_year=2028, user=None):
        self.login(user or self.owner)
        response = self.client.post('/api/ppmp-orders', json=self.draft_payload(fiscal_year))
        self.assertEqual(response.status_code, 201, response.get_json())
        return response.get_json()

    def complete_draft(self, order):
        fields = {
            'project_type': 'Goods/Services', 'procurement_method': 'Small Value Procurement',
            'pre_procurement_conference': 'No', 'procurement_start': '2028-01',
            'procurement_end': '2028-02', 'delivery_period': '2028-03',
            'source_of_funds': 'GAA/101',
        }
        response = self.client.put(f"/api/reports/ppmp-orders/{order['id']}/items", json={
            'ppmp_document_number': 3,
            'items': [{'id': row['id'], 'fields': fields} for row in order['items']],
        })
        self.assertEqual(response.status_code, 200, response.get_json())
        return response.get_json()

    def transition(self, order, user, action, expected=200, **changes):
        self.login(user)
        response = self.client.post(f"/api/ppmp-orders/{order['id']}/transition", json={
            'action': action, **changes,
        })
        self.assertEqual(response.status_code, expected, response.get_json())
        return response.get_json()

    def approved_order(self):
        order = self.create_draft()
        self.complete_draft(order)
        self.transition(order, self.owner, 'submit')
        self.transition(order, self.procurement, 'approve')
        self.transition(order, self.admin, 'approve')
        self.login(self.owner)
        return order

    def create_copy(self, order):
        response = self.client.post(f"/api/ppmp-orders/{order['id']}/indicative-copy", json={})
        self.assertEqual(response.status_code, 201, response.get_json())
        return response.get_json()

    def test_draft_is_unique_per_office_year_and_keeps_custom_items_out_of_master(self):
        first = self.create_draft()
        self.assertEqual(first['status'], 'Draft')
        self.assertTrue(first['workflow_official'])
        registry = OfficialPPMP.query.one()
        self.assertEqual((registry.office_id, registry.fiscal_year, registry.source_order_id),
                         (self.office.id, 2028, first['id']))
        duplicate = self.client.post('/api/ppmp-orders', json=self.draft_payload())
        self.assertEqual(duplicate.status_code, 409)
        self.assertEqual(PPMPOrder.query.count(), 1)
        self.assertEqual(PPMPOrderItem.query.count(), 2)
        self.assertEqual(WishlistItem.query.count(), 1)
        self.create_draft(2029)
        self.assertEqual(OfficialPPMP.query.count(), 2)

    def test_submission_requires_complete_report_and_locks_submitted_data(self):
        order = self.create_draft()
        self.transition(order, self.owner, 'submit', expected=409)
        self.complete_draft(order)
        submitted = self.transition(order, self.owner, 'submit')
        self.assertEqual(submitted['status'], 'Submitted')
        self.assertEqual(submitted['current_stage'], 'Procurement Review')
        response = self.client.put(f"/api/reports/ppmp-orders/{order['id']}/items", json={
            'items': [{'id': order['items'][0]['id'], 'fields': {'remarks': 'Should not save'}}],
        })
        self.assertEqual(response.status_code, 409)
        self.assertIsNone(db.session.get(PPMPOrderItem, order['items'][0]['id']).remarks)
        detail = self.client.get(f"/api/reports/ppmp-orders/{order['id']}").get_json()
        self.assertFalse(detail['editable'])

    def test_returned_correction_resumes_admin_and_approval_locks_original(self):
        order = self.create_draft()
        self.complete_draft(order)
        self.transition(order, self.owner, 'submit')
        pending = self.transition(order, self.procurement, 'approve')
        self.assertEqual(pending['status'], 'Pending Admin Approval')
        returned = self.transition(order, self.admin, 'return', note='Clarify delivery conditions')
        self.assertEqual(returned['return_stage'], 'ADMIN')
        self.login(self.owner)
        response = self.client.put(f"/api/reports/ppmp-orders/{order['id']}/items", json={
            'items': [{'id': order['items'][0]['id'], 'fields': {'remarks': 'Delivery clarified'}}],
        })
        self.assertEqual(response.status_code, 200, response.get_json())
        resumed = self.transition(order, self.owner, 'submit')
        self.assertEqual(resumed['status'], 'Pending Admin Approval')
        approved = self.transition(order, self.admin, 'approve')
        self.assertEqual(approved['status'], 'Approved')
        self.assertEqual(PPMPItem.query.count(), 2)
        self.login(self.owner)
        detail = self.client.get(f"/api/reports/ppmp-orders/{order['id']}").get_json()
        self.assertFalse(detail['editable'])
        self.assertTrue(any(event['details'].get('reviewer_role') == 'ADMIN'
                            for event in detail['history']))
        locked = self.client.put(f"/api/reports/ppmp-orders/{order['id']}/items", json={
            'ppmp_document_number': 99, 'items': [],
        })
        self.assertEqual(locked.status_code, 409)
        self.assertEqual(db.session.get(PPMPOrder, order['id']).ppmp_document_number, 3)

    def test_procurement_return_resumes_procurement_and_eo_cannot_approve_ppmp(self):
        order = self.create_draft()
        self.complete_draft(order)
        self.transition(order, self.owner, 'submit')
        self.transition(order, self.eo, 'approve', expected=409)
        returned = self.transition(order, self.procurement, 'return', note='Check specifications')
        self.assertEqual(returned['return_stage'], 'PROCUREMENT')
        resumed = self.transition(order, self.owner, 'submit')
        self.assertEqual(resumed['status'], 'Submitted')
        self.assertEqual(PPMPItem.query.count(), 0)

    def test_copy_requires_approved_official_and_reopening_preserves_saved_copy(self):
        draft = self.create_draft()
        blocked = self.client.post(f"/api/ppmp-orders/{draft['id']}/indicative-copy", json={})
        self.assertEqual(blocked.status_code, 409)
        self.complete_draft(draft)
        self.transition(draft, self.owner, 'submit')
        self.transition(draft, self.procurement, 'approve')
        self.transition(draft, self.admin, 'approve')
        self.login(self.owner)
        copy = self.create_copy(draft)
        response = self.client.put(copy['open_url'], json={
            'ppmp_document_number': 8,
            'items': [{'id': draft['items'][0]['id'], 'fields': {'remarks': 'Saved copy only'}}],
        })
        self.assertEqual(response.status_code, 200, response.get_json())
        reopened = self.client.post(f"/api/ppmp-orders/{draft['id']}/indicative-copy", json={})
        self.assertEqual(reopened.status_code, 200)
        self.assertEqual(reopened.get_json()['id'], copy['id'])
        self.assertEqual(reopened.get_json()['ppmp_document_number'], 8)
        self.assertEqual(IndicativePPMPCopy.query.count(), 1)

    def test_copy_edits_never_change_approved_source_or_materialized_item(self):
        order = self.approved_order()
        source = db.session.get(PPMPOrderItem, order['items'][0]['id'])
        source_before = deepcopy(source.to_dict())
        materialized = db.session.get(PPMPItem, source.ppmp_item_id)
        materialized_before = deepcopy(materialized.to_dict())
        copy = self.create_copy(order)
        response = self.client.put(copy['open_url'], json={
            'ppmp_document_number': 9,
            'items': [{'id': source.id, 'fields': {
                'general_description': 'Revised indicative description',
                'quantity': '5.50', 'unit': 'box', 'estimated_budget': '987.65',
                'source_of_funds': 'SB 164',
            }}],
        })
        self.assertEqual(response.status_code, 200, response.get_json())
        changed = response.get_json()['items'][0]
        self.assertEqual(changed['general_description'], 'Revised indicative description')
        self.assertEqual(changed['quantity'], 5.5)
        self.assertEqual(changed['estimated_budget'], 987.65)
        db.session.expire_all()
        self.assertEqual(db.session.get(PPMPOrderItem, source.id).to_dict(), source_before)
        self.assertEqual(db.session.get(PPMPItem, materialized.id).to_dict(), materialized_before)
        self.assertEqual(db.session.get(PPMPOrder, order['id']).ppmp_document_number, 3)

    def test_nonowners_cannot_edit_official_or_indicative_copy(self):
        draft = self.create_draft()
        self.login(self.colleague)
        forbidden = self.client.put(f"/api/reports/ppmp-orders/{draft['id']}/items", json={
            'items': [{'id': draft['items'][0]['id'], 'fields': {'quantity': '3'}}],
        })
        self.assertEqual(forbidden.status_code, 409)
        self.login(self.owner)
        self.complete_draft(draft)
        self.transition(draft, self.owner, 'submit')
        self.transition(draft, self.procurement, 'approve')
        self.transition(draft, self.admin, 'approve')
        self.login(self.owner)
        copy = self.create_copy(draft)
        for user in (self.colleague, self.other):
            with self.subTest(role=user.role):
                self.login(user)
                response = self.client.put(copy['open_url'], json={'items': [], 'ppmp_document_number': 44})
                self.assertEqual(response.status_code, 403)
        self.assertEqual(db.session.get(IndicativePPMPCopy, copy['id']).ppmp_document_number, 3)
        self.assertEqual(self.client.get(copy['open_url']).status_code, 403)

    def test_bulk_save_rejects_foreign_item_atomically(self):
        own = self.create_draft()
        foreign = self.create_draft(user=self.other)
        self.login(self.owner)
        before_events = WorkflowAuditLog.query.count()
        response = self.client.put(f"/api/reports/ppmp-orders/{own['id']}/items", json={
            'ppmp_document_number': 7,
            'items': [
                {'id': own['items'][0]['id'], 'fields': {'remarks': 'Must not save'}},
                {'id': foreign['items'][0]['id'], 'fields': {'quantity': '12'}},
            ],
        })
        self.assertEqual(response.status_code, 404)
        self.assertIsNone(db.session.get(PPMPOrderItem, own['items'][0]['id']).remarks)
        self.assertEqual(db.session.get(PPMPOrder, own['id']).ppmp_document_number, 1)
        self.assertEqual(WorkflowAuditLog.query.count(), before_events)

    def test_invalid_report_decimals_are_rejected_without_partial_changes(self):
        order = self.create_draft()
        for invalid in ('NaN', 'Infinity', '-1', '0', '1.001', True):
            with self.subTest(quantity=invalid):
                response = self.client.put(f"/api/reports/ppmp-orders/{order['id']}/items", json={
                    'ppmp_document_number': 5,
                    'items': [
                        {'id': order['items'][0]['id'], 'fields': {'remarks': 'Must roll back'}},
                        {'id': order['items'][1]['id'], 'fields': {'quantity': invalid}},
                    ],
                })
                self.assertEqual(response.status_code, 400, response.get_json())
                self.assertIsNone(db.session.get(PPMPOrderItem, order['items'][0]['id']).remarks)
                self.assertEqual(db.session.get(PPMPOrder, order['id']).ppmp_document_number, 1)

    def test_invalid_new_draft_decimal_does_not_reserve_office_year(self):
        for invalid in ('NaN', 'Infinity', '-1', '0', '1.001', True):
            with self.subTest(quantity=invalid):
                payload = self.draft_payload()
                payload['items'][1]['quantity'] = invalid
                response = self.client.post('/api/ppmp-orders', json=payload)
                self.assertEqual(response.status_code, 400, response.get_json())
                self.assertEqual(PPMPOrder.query.count(), 0)
                self.assertEqual(OfficialPPMP.query.count(), 0)
                self.assertEqual(PPMPOrderItem.query.count(), 0)

    def test_copy_invalid_row_or_decimal_rejects_the_entire_save(self):
        order = self.approved_order()
        copy = self.create_copy(order)
        saved = db.session.get(IndicativePPMPCopy, copy['id'])
        before = deepcopy(saved.items)
        before_events = WorkflowAuditLog.query.count()
        invalid_changes = (
            {'id': order['items'][1]['id'], 'fields': {'estimated_budget': 'NaN'}},
            {'id': 999999, 'fields': {'remarks': 'Unknown row'}},
            {'id': order['items'][1]['id'], 'fields': {'source_order_id': 123}},
            {'id': order['items'][1]['id'], 'fields': {'procurement_end': '2027-01'}},
        )
        for invalid in invalid_changes:
            with self.subTest(invalid=invalid):
                response = self.client.put(copy['open_url'], json={
                    'ppmp_document_number': 29,
                    'items': [
                        {'id': order['items'][0]['id'], 'fields': {'remarks': 'Must not persist'}},
                        invalid,
                    ],
                })
                self.assertEqual(response.status_code, 400, response.get_json())
                db.session.expire_all()
                stored = db.session.get(IndicativePPMPCopy, copy['id'])
                self.assertEqual(stored.items, before)
                self.assertEqual(stored.ppmp_document_number, 3)
                self.assertEqual(WorkflowAuditLog.query.count(), before_events)

    def test_append_remove_follow_owner_status_and_item_membership_rules(self):
        own = self.create_draft()
        foreign = self.create_draft(user=self.other)
        self.login(self.owner)
        path = f"/api/ppmp-orders/{own['id']}/items"
        extra = {'item_name': 'Added requirement', 'category': 'Goods', 'unit': 'piece',
                 'quantity': '3', 'estimated_unit_cost': '10'}
        appended = self.client.post(path, json={'items': [extra]})
        self.assertEqual(appended.status_code, 201, appended.get_json())
        self.assertEqual(len(appended.get_json()['items']), 3)
        duplicate = self.client.post(path, json={'items': [{'master_item_id': self.master.id, 'quantity': '1'}]})
        self.assertEqual(duplicate.status_code, 400)
        self.assertEqual(PPMPOrderItem.query.filter_by(order_id=own['id']).count(), 3)
        removed_foreign = self.client.delete(path + f"/{foreign['items'][0]['id']}")
        self.assertEqual(removed_foreign.status_code, 404)
        added_id = appended.get_json()['items'][-1]['id']
        self.assertEqual(self.client.delete(path + f'/{added_id}').status_code, 200)
        self.login(self.colleague)
        self.assertEqual(self.client.post(path, json={'items': [extra]}).status_code, 403)
        self.assertEqual(self.client.delete(path + f"/{own['items'][0]['id']}").status_code, 403)
        self.login(self.owner)
        self.complete_draft(own)
        self.transition(own, self.owner, 'submit')
        self.assertEqual(self.client.post(path, json={'items': [extra]}).status_code, 403)
        self.assertEqual(self.client.delete(path + f"/{own['items'][0]['id']}").status_code, 403)

    def test_json_false_saves_no_for_preprocurement_conference(self):
        order = self.create_draft()
        response = self.client.put(f"/api/reports/ppmp-orders/{order['id']}/items", json={
            'items': [{'id': row['id'], 'fields': {'pre_procurement_conference': False}}
                      for row in order['items']],
        })
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertTrue(all(row['pre_procurement_conference'] == 'No'
                            for row in response.get_json()['items']))

    def test_admin_cannot_mutate_locked_official_materialization_through_legacy_item_api(self):
        order = self.approved_order()
        source = db.session.get(PPMPOrderItem, order['items'][0]['id'])
        materialized_id = source.ppmp_item_id
        self.login(self.admin)
        for method, payload in (('put', {'quantity': 999, 'authorized_budget': 999}), ('delete', None)):
            with self.subTest(method=method):
                response = getattr(self.client, method)(f'/api/ppmp/{materialized_id}', json=payload)
                self.assertIn(response.status_code, (403, 409), response.get_json())
                db.session.expire_all()
                self.assertIsNotNone(db.session.get(PPMPItem, materialized_id))
                self.assertEqual(db.session.get(PPMPItem, materialized_id).quantity, Decimal('2.00'))

    def test_legacy_designation_preserves_duplicates_and_has_no_requested_dead_end(self):
        legacy = PPMPOrder(
            reference_number='LEGACY-ONE', office_id=self.office.id,
            requested_by=self.owner.id, fiscal_year=2028, status='Requested',
        )
        other_legacy = PPMPOrder(
            reference_number='LEGACY-TWO', office_id=self.office.id,
            requested_by=self.owner.id, fiscal_year=2028, status='Approved',
        )
        db.session.add_all([legacy, other_legacy])
        db.session.commit()
        self.assertEqual(self.client.post(f'/api/ppmp-orders/{legacy.id}/designate-official', json={}).status_code, 403)
        self.login(self.admin)
        designated = self.client.post(f'/api/ppmp-orders/{legacy.id}/designate-official', json={})
        self.assertIn(designated.status_code, (200, 409), designated.get_json())
        if designated.status_code == 200:
            self.assertEqual(designated.get_json()['status'], 'Draft')
            blocked = self.client.post(f'/api/ppmp-orders/{other_legacy.id}/designate-official', json={})
            self.assertEqual(blocked.status_code, 409)
        else:
            self.assertEqual(OfficialPPMP.query.count(), 0)
        self.assertEqual(PPMPOrder.query.count(), 2)
        self.assertEqual(db.session.get(PPMPOrder, other_legacy.id).status, 'Approved')

    def test_official_and_indicative_pdf_excel_exports_use_their_own_data(self):
        order = self.approved_order()
        copy = self.create_copy(order)
        saved = self.client.put(copy['open_url'], json={
            'ppmp_document_number': 11,
            'items': [{'id': order['items'][0]['id'], 'fields': {'general_description': 'COPY ONLY LABEL'}}],
        })
        self.assertEqual(saved.status_code, 200)
        for base, number, copy_expected in (
            (f"/api/reports/ppmp-orders/{order['id']}/export", 3, False),
            (f"/api/ppmp-indicative-copies/{copy['id']}/export", 11, True),
        ):
            with self.subTest(base=base):
                pdf = self.client.get(base + '/pdf')
                self.assertEqual(pdf.status_code, 200)
                self.assertEqual(pdf.mimetype, 'application/pdf')
                self.assertTrue(pdf.data.startswith(b'%PDF-'))
                self.assertGreater(len(pdf.data), 1000)
                excel = self.client.get(base + '/excel')
                self.assertEqual(excel.status_code, 200)
                workbook = load_workbook(io.BytesIO(excel.data), data_only=True)
                sheet = workbook.active
                self.assertIn(f'NO. {number}', sheet['A1'].value)
                values = [value for row in sheet.iter_rows(values_only=True) for value in row]
                self.assertEqual('COPY ONLY LABEL' in values, copy_expected)
                workbook.close()


class AdditiveWorkflowMigrationTest(unittest.TestCase):
    def test_legacy_duplicates_and_workflow_versions_are_preserved_on_repeated_migration(self):
        migration_app = Flask('isolated-workflow-migration')
        migration_app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///:memory:'
        db.init_app(migration_app)
        with migration_app.app_context():
            # Model only the columns that existed before the new workflow.
            # Historical duplicate orders must survive the new registry rule.
            definitions = {
                'colleges': 'id INTEGER PRIMARY KEY', 'users': 'id INTEGER PRIMARY KEY',
                'ppmp_orders': 'id INTEGER PRIMARY KEY, office_id INTEGER, fiscal_year INTEGER',
                'ppmp_records': 'id INTEGER PRIMARY KEY',
                'annual_procurement_plans': 'id INTEGER PRIMARY KEY',
                'ppmp_items': 'id INTEGER PRIMARY KEY, mode_of_procurement VARCHAR(120)',
                'app_items': 'id INTEGER PRIMARY KEY, mode_of_procurement VARCHAR(180)',
                'app_approvals': 'id INTEGER PRIMARY KEY',
                'ppmp_record_items': 'id INTEGER PRIMARY KEY, procurement_method VARCHAR(120)',
                'consolidation_batches': 'id INTEGER PRIMARY KEY',
            }
            try:
                with db.engine.begin() as connection:
                    for table, definition in definitions.items():
                        connection.execute(text(f'CREATE TABLE {table} ({definition})'))
                    connection.execute(text('INSERT INTO colleges (id) VALUES (1)'))
                    connection.execute(text(
                        'INSERT INTO ppmp_orders (id, office_id, fiscal_year) '
                        'VALUES (1, 1, 2028), (2, 1, 2028)'
                    ))
                    connection.execute(text('INSERT INTO annual_procurement_plans (id) VALUES (1)'))
                first = apply_workflow_migrations(db)
                self.assertEqual(len(first['added_columns']), sum(map(len, WORKFLOW_REQUIRED_COLUMNS.values())))
                actual_tables = set(inspect(db.engine).get_table_names())
                self.assertTrue(set(WORKFLOW_REQUIRED_TABLES).issubset(actual_tables))
                with db.engine.connect() as connection:
                    self.assertEqual(connection.execute(text(
                        'SELECT id, office_id, fiscal_year, workflow_official FROM ppmp_orders ORDER BY id'
                    )).all(), [(1, 1, 2028, 0), (2, 1, 2028, 0)])
                    self.assertEqual(connection.execute(text(
                        'SELECT workflow_version, revision FROM annual_procurement_plans WHERE id=1'
                    )).one(), (1, 1))
                    self.assertEqual(connection.execute(text('SELECT COUNT(*) FROM official_ppmps')).scalar_one(), 0)
                repeated = apply_workflow_migrations(db)
                self.assertEqual(repeated, {'created_tables': [], 'added_columns': [], 'widened_columns': []})
            finally:
                db.session.remove()
                db.engine.dispose()


if __name__ == '__main__':
    unittest.main()
